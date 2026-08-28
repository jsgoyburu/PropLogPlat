import datetime
import json
import zipfile
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils.translation import override as language_override
try:
    from openpyxl import Workbook, load_workbook
except ModuleNotFoundError:  # pragma: no cover - depende del entorno
    Workbook = None
    load_workbook = None

from accounts.models import EncuestaEstudiante
from cursos.models import Cohorte, Comision, Inscripcion, Parcial, NotaParcial
from ejercicios.models import Ejercicio, EjercicioPractica, Intento, Practica, PracticaComision, Progreso
from .forms import EjercicioForm


Usuario = get_user_model()


def _u(username, password='clave123', **kwargs):
    """Crea usuario con consentimientos completados para pasar el middleware."""
    user = Usuario.objects.create_user(username=username, password=password, **kwargs)
    Usuario.objects.filter(pk=user.pk).update(
        consentimiento_pedagogico=True,
        consentimiento_investigacion=True,
        encuesta_completada=True,
    )
    user.refresh_from_db()
    return user


class EjercicioFormTests(TestCase):
    def test_rechaza_formula_solucion_invalida(self):
        form = EjercicioForm(data={
            'enunciado': 'Probar una fórmula inválida',
            'formula_solucion': 'p && q',
            'tipo': 'formalizacion',
            'es_publico': False,
        })

        self.assertFalse(form.is_valid())
        self.assertIn('formula_solucion', form.errors)

    def test_normaliza_formula_solucion_antes_de_guardar(self):
        form = EjercicioForm(data={
            'enunciado': 'Probar normalización',
            'formula_solucion': 'p ∧ q',
            'tipo': 'formalizacion',
            'es_publico': True,
        })

        self.assertTrue(form.is_valid())
        self.assertEqual(form.cleaned_data['formula_solucion'], 'p · q')


class PaquetesPortablesTests(TestCase):
    def setUp(self):
        self.docente = _u('doc-paquetes', es_docente=True)
        self.client.login(username='doc-paquetes', password='clave123')
        self.practica = Practica.objects.create(
            titulo='Práctica compartida',
            titulo_en='Shared practice',
            titulo_zh_hans='共享练习',
            descripcion='<p>Descripción <strong>formateada</strong></p>',
            creada_por=self.docente,
        )
        self.ejercicio = Ejercicio.objects.create(
            enunciado='<p>Si llueve, la calle se moja.</p>',
            enunciado_en='If it rains, the street gets wet.',
            enunciado_fr='S’il pleut, la rue est mouillée.',
            enunciado_de='Wenn es regnet, wird die Straße nass.',
            enunciado_zh_hans='如果下雨，街道就会湿。',
            formula_solucion='p -> q',
            tipo='formalizacion',
            diccionario_solucion={'p': 'Llueve', 'q': 'La calle se moja'},
            creado_por=self.docente,
        )
        EjercicioPractica.objects.create(
            practica=self.practica,
            ejercicio=self.ejercicio,
            orden=1,
        )

    def _descargar_practica(self):
        response = self.client.get(
            reverse('docentes:practica_descargar', args=[self.practica.id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/zip')
        return response.content

    def test_descarga_zip_json_sin_datos_personales(self):
        contenido = self._descargar_practica()
        with zipfile.ZipFile(BytesIO(contenido)) as archivo:
            self.assertEqual(archivo.namelist(), ['package.json'])
            payload = json.loads(archivo.read('package.json'))

        self.assertEqual(payload['format'], 'ipc-logica-package')
        self.assertEqual(payload['version'], 1)
        self.assertEqual(payload['kind'], 'practice')
        self.assertEqual(payload['content']['title']['en'], 'Shared practice')
        self.assertEqual(payload['content']['title']['zh-hans'], '共享练习')
        self.assertEqual(
            payload['content']['exercises'][0]['enunciado']['de'],
            'Wenn es regnet, wird die Straße nass.',
        )
        self.assertNotIn('creado_por', json.dumps(payload))
        self.assertNotIn('doc-paquetes', json.dumps(payload))
        self.assertNotIn('<strong>', json.dumps(payload))

    def test_instala_practica_como_copia_privada(self):
        contenido = self._descargar_practica()
        subida = SimpleUploadedFile(
            'practica.zip', contenido, content_type='application/zip',
        )
        response = self.client.post(
            reverse('docentes:paquete_instalar'),
            {'archivo': subida},
        )

        self.assertEqual(response.status_code, 302)
        copia = Practica.objects.exclude(pk=self.practica.pk).get()
        self.assertEqual(copia.titulo, 'Práctica compartida')
        self.assertEqual(copia.titulo_en, 'Shared practice')
        self.assertEqual(copia.titulo_zh_hans, '共享练习')
        self.assertFalse(copia.es_publica)
        self.assertEqual(copia.creada_por, self.docente)
        self.assertIsNone(copia.practica_origen)
        ejercicio_copia = copia.ejercicio_practicas.get().ejercicio
        self.assertFalse(ejercicio_copia.es_publico)
        self.assertEqual(ejercicio_copia.creado_por, self.docente)
        self.assertEqual(ejercicio_copia.enunciado_fr, 'S’il pleut, la rue est mouillée.')
        self.assertEqual(ejercicio_copia.enunciado_zh_hans, '如果下雨，街道就会湿。')

    def test_chino_localizado_y_fallback_al_castellano(self):
        with language_override('zh-hans'):
            self.assertEqual(self.practica.titulo_localizado, '共享练习')
            self.assertEqual(self.ejercicio.enunciado_localizado, '如果下雨，街道就会湿。')

            self.practica.titulo_zh_hans = ''
            self.ejercicio.enunciado_zh_hans = ''
            self.assertEqual(self.practica.titulo_localizado, 'Práctica compartida')
            self.assertIn('Si llueve', self.ejercicio.enunciado_localizado)

    def test_instala_paquete_v1_anterior_sin_clave_china(self):
        contenido = self._descargar_practica()
        with zipfile.ZipFile(BytesIO(contenido)) as archivo:
            payload = json.loads(archivo.read('package.json'))
        payload['content']['title'].pop('zh-hans')
        payload['content']['description'].pop('zh-hans')
        for ejercicio in payload['content']['exercises']:
            ejercicio['enunciado'].pop('zh-hans')

        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED) as archivo:
            archivo.writestr('package.json', json.dumps(payload, ensure_ascii=False))
        response = self.client.post(
            reverse('docentes:paquete_instalar'),
            {'archivo': SimpleUploadedFile('v1-anterior.zip', buffer.getvalue())},
        )

        self.assertEqual(response.status_code, 302)
        copia = Practica.objects.exclude(pk=self.practica.pk).get()
        self.assertEqual(copia.titulo_zh_hans, '')
        self.assertEqual(copia.ejercicio_practicas.get().ejercicio.enunciado_zh_hans, '')

    def test_rechaza_zip_con_archivos_adicionales(self):
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archivo:
            archivo.writestr('package.json', '{}')
            archivo.writestr('extra.txt', 'no')
        subida = SimpleUploadedFile('malo.zip', buffer.getvalue())

        response = self.client.post(
            reverse('docentes:paquete_instalar'),
            {'archivo': subida},
            follow=True,
        )

        self.assertContains(response, 'únicamente un archivo package.json')
        self.assertEqual(Practica.objects.count(), 1)

    def test_descarga_ejercicio_individual(self):
        response = self.client.get(
            reverse('docentes:ejercicio_descargar', args=[self.ejercicio.id])
        )
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(BytesIO(response.content)) as archivo:
            payload = json.loads(archivo.read('package.json'))
        self.assertEqual(payload['kind'], 'exercise')
        self.assertEqual(payload['content']['enunciado']['es'], 'Si llueve, la calle se moja.')


class GestionEstudiantesTests(TestCase):
    def setUp(self):
        self.docente = _u('docente', es_docente=True)
        self.docente_otro = _u('docente2', es_docente=True)
        self.admin = _u('admin', is_staff=True, is_superuser=True)
        self.comision = Comision.objects.create(nombre='IPC Noche')
        self.comision.docentes.add(self.docente)

    def _crear_archivo_excel(self, filas):
        if Workbook is None:
            self.skipTest('openpyxl no está instalado en el entorno de tests.')
        wb = Workbook()
        ws = wb.active
        ws.append(['username', 'email', 'password'])
        for fila in filas:
            ws.append(fila)
        archivo = BytesIO()
        wb.save(archivo)
        archivo.seek(0)
        archivo.name = 'estudiantes.xlsx'
        return archivo

    def test_alta_individual_crea_estudiante_e_inscripcion(self):
        self.client.login(username='docente', password='clave123')
        response = self.client.post(
            reverse('docentes:estudiante_create', args=[self.comision.id]),
            data={'username': 'alu1', 'email': 'alu1@example.com', 'password': 'pass1234'},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        estudiante = Usuario.objects.get(username='alu1')
        self.assertFalse(estudiante.es_docente)
        self.assertFalse(estudiante.is_staff)
        self.assertTrue(Inscripcion.objects.filter(estudiante=estudiante, comision=self.comision).exists())

    def test_importacion_omite_filas_invalidas_y_reporta(self):
        Usuario.objects.create_user(username='repetido', email='dup@example.com', password='x')
        self.client.login(username='docente', password='clave123')
        archivo = self._crear_archivo_excel([
            ['nuevo1', 'nuevo1@example.com', 'abc12345'],
            ['repetido', 'otro@example.com', 'abc12345'],
            ['nuevo2', 'mail-invalido', 'abc12345'],
            ['nuevo3', 'nuevo3@example.com', ''],
        ])

        response = self.client.post(
            reverse('docentes:estudiantes_importar', args=[self.comision.id]),
            data={'archivo': archivo},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'procesadas=4')
        self.assertContains(response, 'importadas=1')
        self.assertContains(response, 'fallidas=3')
        self.assertTrue(Usuario.objects.filter(username='nuevo1').exists())
        self.assertFalse(Usuario.objects.filter(username='nuevo2').exists())
        self.assertContains(response, 'Fila 3')

    def test_importacion_restringida_a_docente_de_comision_o_admin(self):
        archivo = self._crear_archivo_excel([['nuevox', 'nuevox@example.com', 'abc12345']])

        self.client.login(username='docente2', password='clave123')
        response = self.client.post(
            reverse('docentes:estudiantes_importar', args=[self.comision.id]),
            data={'archivo': archivo},
        )
        self.assertEqual(response.status_code, 403)

        archivo_admin = self._crear_archivo_excel([['nuevoa', 'nuevoa@example.com', 'abc12345']])
        self.client.login(username='admin', password='clave123')
        response_admin = self.client.post(
            reverse('docentes:estudiantes_importar', args=[self.comision.id]),
            data={'archivo': archivo_admin},
        )
        self.assertEqual(response_admin.status_code, 200)
        self.assertTrue(Usuario.objects.filter(username='nuevoa').exists())

    def test_alta_individual_sin_cohorte_activa_no_crea_nada(self):
        """Sin cohorte activa no hay cohorte que reciba el alta: la guarda de
        cohorte cerrada (`_require_cohorte_activa`) corta con 403, no con un
        IntegrityError ni con una mensaje-y-redirect silencioso."""
        Cohorte.objects.filter(activa=True).update(activa=False)
        self.client.login(username='docente', password='clave123')

        response = self.client.post(
            reverse('docentes:estudiante_create', args=[self.comision.id]),
            data={'username': 'alu_sin_cohorte', 'email': 'alusc@example.com', 'password': 'pass1234'},
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Usuario.objects.filter(username='alu_sin_cohorte').exists())

    def test_importacion_sin_cohorte_activa_no_crea_nada(self):
        """La importación masiva tampoco puede escribir sin cohorte activa:
        misma guarda de cohorte cerrada que el alta individual."""
        Cohorte.objects.filter(activa=True).update(activa=False)
        self.client.login(username='docente', password='clave123')
        archivo = self._crear_archivo_excel([['alu_import_sc', 'aisc@example.com', 'abc12345']])

        response = self.client.post(
            reverse('docentes:estudiantes_importar', args=[self.comision.id]),
            data={'archivo': archivo},
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Usuario.objects.filter(username='alu_import_sc').exists())

    def test_exportacion_excel_estudiantes_desde_comision(self):
        if load_workbook is None:
            self.skipTest('openpyxl no está instalado en el entorno de tests.')
        estudiante = _u(
            'alu-export',
            first_name='Ada',
            last_name='Lovelace',
            email='ada@example.com',
        )
        Inscripcion.objects.create(
            estudiante=estudiante, comision=self.comision,
            cohorte=Cohorte.objects.get(anio=2026, cuatrimestre=1),
        )
        encuesta, _ = EncuestaEstudiante.objects.get_or_create(estudiante=estudiante)
        encuesta.dni = '12345678'
        encuesta.save(update_fields=['dni'])

        self.client.login(username='docente', password='clave123')
        response = self.client.get(reverse('docentes:estudiantes_exportar', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment; filename=', response['Content-Disposition'])
        wb = load_workbook(filename=BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, 'DNI')
        self.assertEqual(ws.cell(row=1, column=2).value, 'Apellido')
        self.assertEqual(ws.cell(row=1, column=3).value, 'Nombre')
        self.assertEqual(ws.cell(row=1, column=4).value, 'Correo electrónico')
        self.assertEqual(ws.cell(row=2, column=1).value, '12345678')
        self.assertEqual(ws.cell(row=2, column=2).value, 'Lovelace')
        self.assertEqual(ws.cell(row=2, column=3).value, 'Ada')
        self.assertEqual(ws.cell(row=2, column=4).value, 'ada@example.com')

    def test_exportacion_no_mezcla_camadas_de_un_recursante(self):
        """Code review (Important I5): sin acotar por cohorte, un recursante
        con dos Inscripcion en la comisión (una por cohorte) aparecía dos
        veces en el XLSX y mezclaba datos de camadas distintas. El botón del
        panel debe filtrar igual que exportar_cohorte (management command,
        ya corregido) y que el resto del panel: la cohorte vista, o la activa
        por default."""
        if load_workbook is None:
            self.skipTest('openpyxl no está instalado en el entorno de tests.')
        cohorte_vieja = Cohorte.objects.create(anio=2020, cuatrimestre=1)
        cohorte_actual = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        recursante = _u('alu-export-recursa', first_name='Grace', last_name='Hopper')
        Inscripcion.objects.create(estudiante=recursante, comision=self.comision, cohorte=cohorte_vieja)
        Inscripcion.objects.create(estudiante=recursante, comision=self.comision, cohorte=cohorte_actual)

        self.client.login(username='docente', password='clave123')
        response = self.client.get(reverse('docentes:estudiantes_exportar', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        wb = load_workbook(filename=BytesIO(response.content))
        ws = wb.active
        apellidos = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(apellidos.count('Hopper'), 1)

    def test_exportacion_respeta_el_parametro_cohorte_explicito(self):
        """Contraparte: con ?cohorte=<vieja>, exporta la camada vieja, no la activa."""
        if load_workbook is None:
            self.skipTest('openpyxl no está instalado en el entorno de tests.')
        cohorte_vieja = Cohorte.objects.create(anio=2020, cuatrimestre=1)
        cohorte_actual = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        recursante = _u('alu-export-recursa2', first_name='Katherine', last_name='Johnson')
        Inscripcion.objects.create(estudiante=recursante, comision=self.comision, cohorte=cohorte_vieja)
        Inscripcion.objects.create(estudiante=recursante, comision=self.comision, cohorte=cohorte_actual)

        self.client.login(username='docente', password='clave123')
        url = reverse('docentes:estudiantes_exportar', args=[self.comision.id])
        response = self.client.get(f'{url}?cohorte={cohorte_vieja.pk}')

        self.assertEqual(response.status_code, 200)
        wb = load_workbook(filename=BytesIO(response.content))
        ws = wb.active
        apellidos = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(apellidos.count('Johnson'), 1)


class EstudianteEditRemoveRecursanteTests(TestCase):
    """Code review (Critical C2 y C3): un recursante tiene dos Inscripcion en
    la misma comisión, una por cohorte.

    C2: estudiante_edit/estudiante_remove resolvían el estudiante con un
    get_object_or_404 que hacía join sobre inscripciones__comision; para un
    recursante ese join devuelve dos filas y explota con
    Usuario.MultipleObjectsReturned (500).

    C3: una vez resuelto el 500, estudiante_remove borraba TODAS las
    Inscripcion de esa comisión sin filtrar por cohorte, así que dar de baja
    a alguien de la camada actual también borraba su inscripción de la
    camada anterior (el único registro de que perteneció a esa camada)."""

    def setUp(self):
        self.docente = _u('doc-edit-recursa', es_docente=True)
        self.estudiante = _u('est-recursa-edit', first_name='Ada', last_name='Lovelace')

        # 2026-C1 es la activa del backfill; se usa como "cohorte actual" y se
        # crea aparte la "vieja" para el recursante.
        self.cohorte_actual = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.cohorte_vieja = Cohorte.objects.create(anio=2020, cuatrimestre=1)

        self.comision = Comision.objects.create(nombre='IPC Recursa Edit')
        self.comision.docentes.add(self.docente)

        # Vieja primero, actual después: mismo orden que el resto de la
        # suite usa para setear la inscripción "más reciente".
        self.insc_vieja = Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte_vieja,
        )
        self.insc_actual = Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte_actual,
        )

        self.client.login(username='doc-edit-recursa', password='clave123')

    def test_estudiante_edit_no_revienta_con_recursante(self):
        response = self.client.get(
            reverse('docentes:estudiante_edit', args=[self.comision.id, self.estudiante.id])
        )
        self.assertEqual(response.status_code, 200)

    def test_estudiante_edit_post_no_revienta_con_recursante(self):
        response = self.client.post(
            reverse('docentes:estudiante_edit', args=[self.comision.id, self.estudiante.id]),
            data={
                'username': self.estudiante.username,
                'email': 'ada-edit@example.com',
                'first_name': 'Ada',
                'last_name': 'Lovelace',
            },
        )
        self.assertEqual(response.status_code, 302)
        self.estudiante.refresh_from_db()
        self.assertEqual(self.estudiante.email, 'ada-edit@example.com')

    def test_estudiante_remove_no_revienta_con_recursante(self):
        url = reverse('docentes:estudiante_remove', args=[self.comision.id, self.estudiante.id])
        response = self.client.post(f'{url}?cohorte={self.cohorte_vieja.pk}')
        self.assertEqual(response.status_code, 302)

    def test_estudiante_remove_solo_borra_la_inscripcion_de_la_cohorte_vista(self):
        """La baja actúa sobre la cohorte que el docente está viendo (?cohorte=
        en el querystring), no sobre todas las inscripciones del estudiante en
        la comisión."""
        url = reverse('docentes:estudiante_remove', args=[self.comision.id, self.estudiante.id])
        self.client.post(f'{url}?cohorte={self.cohorte_vieja.pk}')

        self.assertFalse(Inscripcion.objects.filter(pk=self.insc_vieja.pk).exists())
        self.assertTrue(Inscripcion.objects.filter(pk=self.insc_actual.pk).exists())

    def test_estudiante_remove_sin_parametro_cohorte_usa_la_activa(self):
        """Sin ?cohorte= explícito, _resolver_cohorte cae a la activa (misma
        regla que el resto del panel): borra la inscripción de la camada en
        curso, no la vieja."""
        url = reverse('docentes:estudiante_remove', args=[self.comision.id, self.estudiante.id])
        self.client.post(url)

        self.assertTrue(Inscripcion.objects.filter(pk=self.insc_vieja.pk).exists())
        self.assertFalse(Inscripcion.objects.filter(pk=self.insc_actual.pk).exists())

    def test_estudiante_remove_no_borra_progreso_ni_intentos_de_la_cohorte(self):
        """Decisión sobre el efecto dominó: Progreso e Intento no cuelgan de
        Inscripcion (tienen su propio FK a cohorte), así que dar de baja la
        inscripción de una cohorte no debe arrastrar su historial de
        intentos/progreso. Se verifica explícitamente para que un cambio futuro
        que agregue un cascade no lo haga sin que este test lo note."""
        practica = Practica.objects.create(titulo='P-recursa-remove', creada_por=self.docente)
        pc = PracticaComision.objects.create(practica=practica, comision=self.comision, orden=1)
        ejercicio = Ejercicio.objects.create(
            enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        ep = EjercicioPractica.objects.create(practica=practica, ejercicio=ejercicio, orden=1)
        Intento.objects.create(
            estudiante=self.estudiante, ejercicio_practica=ep, practica_comision=pc,
            cohorte=self.cohorte_vieja, respuesta_raw='p', es_correcto=True,
        )
        progreso = Progreso.objects.create(
            estudiante=self.estudiante, practica_comision=pc, cohorte=self.cohorte_vieja,
            ejercicio_practica_actual=None,
        )

        url = reverse('docentes:estudiante_remove', args=[self.comision.id, self.estudiante.id])
        self.client.post(f'{url}?cohorte={self.cohorte_vieja.pk}')

        self.assertTrue(Intento.objects.filter(estudiante=self.estudiante, cohorte=self.cohorte_vieja).exists())
        self.assertTrue(Progreso.objects.filter(pk=progreso.pk).exists())


class ComisionesListViewTests(TestCase):
    def setUp(self):
        self.docente = _u('docente-list', es_docente=True)
        self.docente_otro = _u('docente-otro-list', es_docente=True)
        self.admin = _u('admin-list', is_staff=True, is_superuser=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)

        self.comision_docente = Comision.objects.create(nombre='Comisión Docente')
        self.comision_docente.docentes.add(self.docente)

        self.comision_otra = Comision.objects.create(nombre='Comisión Otra')
        self.comision_otra.docentes.add(self.docente_otro)

    def test_docente_solo_ve_sus_comisiones(self):
        self.client.login(username='docente-list', password='clave123')

        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Comisión Docente')
        self.assertNotContains(response, 'Comisión Otra')

    def test_admin_no_tiene_excepcion_ve_solo_sus_comisiones(self):
        """El staff no ve todas las comisiones: rige el mismo filtro que un docente.

        Comportamiento deliberado desde 7040945 ("Elimina excepción is_staff en
        vistas del panel docente"): los staff ven y operan únicamente sobre las
        comisiones en las que son docentes.
        """
        self.comision_otra.docentes.add(self.admin)
        self.client.login(username='admin-list', password='clave123')

        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Comisión Otra')
        self.assertNotContains(response, 'Comisión Docente')

    def test_muestra_metricas_operativas_y_avance_promedio(self):
        estudiante_1 = _u('est-1')
        estudiante_2 = _u('est-2')
        Inscripcion.objects.create(estudiante=estudiante_1, comision=self.comision_docente, cohorte=self.cohorte)
        Inscripcion.objects.create(estudiante=estudiante_2, comision=self.comision_docente, cohorte=self.cohorte)

        practica = Practica.objects.create(titulo='Práctica 1', creada_por=self.docente)
        pc = PracticaComision.objects.create(practica=practica, comision=self.comision_docente, orden=1)
        ejercicio_1 = Ejercicio.objects.create(enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente)
        ejercicio_2 = Ejercicio.objects.create(enunciado='E2', formula_solucion='q', tipo='formalizacion', creado_por=self.docente)
        ep1 = EjercicioPractica.objects.create(practica=practica, ejercicio=ejercicio_1, orden=1)
        ep2 = EjercicioPractica.objects.create(practica=practica, ejercicio=ejercicio_2, orden=2)

        # Avance docente: estudiante_1 tiene ep1+ep2 aprobados (100%), estudiante_2 solo ep1 (50%)
        # Promedio = (100 + 50) / 2 = 75%
        Intento.objects.create(estudiante=estudiante_1, ejercicio_practica=ep1, practica_comision=pc, cohorte=self.cohorte, respuesta_raw='p', es_correcto=True, aprobado_docente=True)
        Intento.objects.create(estudiante=estudiante_1, ejercicio_practica=ep2, practica_comision=pc, cohorte=self.cohorte, respuesta_raw='q', es_correcto=True, aprobado_docente=True)
        Intento.objects.create(estudiante=estudiante_2, ejercicio_practica=ep1, practica_comision=pc, cohorte=self.cohorte, respuesta_raw='p', es_correcto=True, aprobado_docente=True)

        self.client.login(username='docente-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '2')
        self.assertContains(response, '1')
        self.assertContains(response, '75,0%')

    def test_muestra_avance_resuelto_sin_correccion_manual(self):
        estudiante = _u('est-sin-revisar')
        Inscripcion.objects.create(estudiante=estudiante, comision=self.comision_docente, cohorte=self.cohorte)

        practica = Practica.objects.create(titulo='Práctica 1', creada_por=self.docente)
        pc = PracticaComision.objects.create(
            practica=practica, comision=self.comision_docente, orden=1,
        )
        ejercicio_1 = Ejercicio.objects.create(
            enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        ejercicio_2 = Ejercicio.objects.create(
            enunciado='E2', formula_solucion='q', tipo='formalizacion', creado_por=self.docente,
        )
        ep1 = EjercicioPractica.objects.create(practica=practica, ejercicio=ejercicio_1, orden=1)
        EjercicioPractica.objects.create(practica=practica, ejercicio=ejercicio_2, orden=2)

        # Correcto pero sin revisar: cuenta como resuelto, no como revisado.
        Intento.objects.create(
            estudiante=estudiante, ejercicio_practica=ep1, practica_comision=pc,
            cohorte=self.cohorte, respuesta_raw='p', es_correcto=True, aprobado_docente=None,
        )

        self.client.login(username='docente-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Resuelto')
        self.assertContains(response, 'Revisado')

        # Igual que en el detalle: los porcentajes se verifican en el contexto
        # porque '0,0%' es substring de '50,0%'.
        item = next(
            i for i in response.context['comisiones_data']
            if i['comision'].id == self.comision_docente.id
        )
        self.assertEqual(item['avance_resuelto'], 50.0)
        self.assertEqual(item['avance_promedio'], 0.0)

    def test_recursante_se_cuenta_una_sola_vez(self):
        """Code review (Important I3): cantidad_estudiantes sumaba todas las
        camadas (M2M sin filtrar) y contaba doble a un recursante; el
        promedio de avance pesaba doble ese caso y cruzaba camadas. La
        pantalla debe mostrar la camada en curso (la cohorte activa), igual
        que comision_detail sin ?cohorte= explícito."""
        cohorte_vieja = Cohorte.objects.create(anio=2020, cuatrimestre=1)
        recursante = _u('est-recursa-list')
        Inscripcion.objects.create(
            estudiante=recursante, comision=self.comision_docente, cohorte=cohorte_vieja,
        )
        Inscripcion.objects.create(
            estudiante=recursante, comision=self.comision_docente, cohorte=self.cohorte,
        )

        practica = Practica.objects.create(titulo='Práctica Recursa', creada_por=self.docente)
        pc = PracticaComision.objects.create(practica=practica, comision=self.comision_docente, orden=1)
        ejercicio = Ejercicio.objects.create(
            enunciado='E', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        EjercicioPractica.objects.create(practica=practica, ejercicio=ejercicio, orden=1)

        # Completó SOLO en la camada vieja: en la actual no hizo nada todavía.
        Intento.objects.create(
            estudiante=recursante, ejercicio_practica=EjercicioPractica.objects.get(practica=practica),
            practica_comision=pc, cohorte=cohorte_vieja,
            respuesta_raw='p', es_correcto=True, aprobado_docente=True,
        )

        self.client.login(username='docente-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        item = next(
            i for i in response.context['comisiones_data']
            if i['comision'].id == self.comision_docente.id
        )
        self.assertEqual(item['cantidad_estudiantes'], 1)
        # En la camada actual no resolvió nada: 0%, no lo completado en la vieja.
        self.assertEqual(item['avance_promedio'], 0.0)
        self.assertEqual(item['avance_resuelto'], 0.0)

    def _crear_practicas_de_banco(self):
        """Una práctica propia, una pública ajena y una privada ajena."""
        propia = Practica.objects.create(
            titulo='Práctica Propia', creada_por=self.docente,
        )
        publica_ajena = Practica.objects.create(
            titulo='Práctica Pública Ajena', creada_por=self.docente_otro,
            es_publica=True,
        )
        privada_ajena = Practica.objects.create(
            titulo='Práctica Privada Ajena', creada_por=self.docente_otro,
            es_publica=False,
        )
        return propia, publica_ajena, privada_ajena

    def test_banco_muestra_propias_y_publicas(self):
        """El banco se define por autoría y es_publica, no por comisión."""
        propia, publica_ajena, privada_ajena = self._crear_practicas_de_banco()

        self.client.login(username='docente-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertIn(propia, response.context['mis_practicas'])
        self.assertIn(publica_ajena, response.context['mis_practicas'])
        self.assertNotIn(privada_ajena, response.context['mis_practicas'])

    def test_admin_no_ve_practicas_privadas_ajenas_en_el_banco(self):
        """El banco no tiene excepción is_staff: mismo criterio para todos los roles.

        Es el mismo principio que 7040945 aplicó a las comisiones. El admin conserva
        el permiso de editar/eliminar cualquier práctica; lo que no tiene es un
        listado global en el panel.
        """
        _, publica_ajena, privada_ajena = self._crear_practicas_de_banco()

        self.client.login(username='admin-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertIn(publica_ajena, response.context['mis_practicas'])
        self.assertNotIn(privada_ajena, response.context['mis_practicas'])

    def test_practica_publica_ajena_no_ofrece_agregar_ejercicio(self):
        """Sobre una práctica ajena hay lectura pero no escritura.

        ejercicio_create descarta la preselección con _puede_editar_practica, así que
        ofrecer "+ Ejercicio" terminaría creando un ejercicio sin asignar.
        """
        propia, publica_ajena, _ = self._crear_practicas_de_banco()

        self.client.login(username='docente-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'?practica={propia.id}&next=')
        self.assertNotContains(response, f'?practica={publica_ajena.id}&next=')

    def _crear_ejercicios_de_banco(self):
        """Un ejercicio propio, uno público ajeno y uno privado ajeno."""
        propio = Ejercicio.objects.create(
            enunciado='Ejercicio Propio', formula_solucion='p', tipo='formalizacion',
            creado_por=self.docente,
        )
        publico_ajeno = Ejercicio.objects.create(
            enunciado='Ejercicio Público Ajeno', formula_solucion='q',
            tipo='formalizacion', creado_por=self.docente_otro, es_publico=True,
        )
        privado_ajeno = Ejercicio.objects.create(
            enunciado='Ejercicio Privado Ajeno', formula_solucion='r',
            tipo='formalizacion', creado_por=self.docente_otro, es_publico=False,
        )
        return propio, publico_ajeno, privado_ajeno

    def test_banco_de_ejercicios_muestra_propios_y_publicos(self):
        propio, publico_ajeno, privado_ajeno = self._crear_ejercicios_de_banco()

        self.client.login(username='docente-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertIn(propio, response.context['mis_ejercicios'])
        self.assertIn(publico_ajeno, response.context['mis_ejercicios'])
        self.assertNotIn(privado_ajeno, response.context['mis_ejercicios'])

    def test_admin_no_ve_ejercicios_privados_ajenos_en_el_banco(self):
        _, publico_ajeno, privado_ajeno = self._crear_ejercicios_de_banco()

        self.client.login(username='admin-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertIn(publico_ajeno, response.context['mis_ejercicios'])
        self.assertNotIn(privado_ajeno, response.context['mis_ejercicios'])

    def test_ejercicio_publico_ajeno_no_ofrece_editar_ni_eliminar(self):
        """ejercicio_edit exige ser el autor: la plantilla no debe ofrecer el botón."""
        propio, publico_ajeno, _ = self._crear_ejercicios_de_banco()

        self.client.login(username='docente-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse('docentes:ejercicio_edit', args=[propio.id]))
        self.assertNotContains(
            response, reverse('docentes:ejercicio_edit', args=[publico_ajeno.id]),
        )
        self.assertNotContains(
            response, reverse('docentes:ejercicio_delete', args=[publico_ajeno.id]),
        )

    def test_banco_del_panel_excluye_practicas_derivadas(self):
        """Las copias de trabajo quedan fuera aunque sean propias o públicas."""
        derivada_propia = Practica.objects.create(
            titulo='Copia de trabajo propia', creada_por=self.docente,
            practica_origen=Practica.objects.create(
                titulo='Origen', creada_por=self.docente,
            ),
        )
        derivada_publica = Practica.objects.create(
            titulo='Copia de trabajo pública', creada_por=self.docente_otro,
            es_publica=True,
            practica_origen=Practica.objects.create(
                titulo='Origen ajeno', creada_por=self.docente_otro, es_publica=True,
            ),
        )

        self.client.login(username='docente-list', password='clave123')
        response = self.client.get(reverse('docentes:comisiones_list'))

        self.assertEqual(response.status_code, 200)
        self.assertNotIn(derivada_propia, response.context['mis_practicas'])
        self.assertNotIn(derivada_publica, response.context['mis_practicas'])


class ComisionDetailProgressViewTests(TestCase):
    def setUp(self):
        # comision_detail cachea las analíticas 5 min bajo la clave
        # `analiticas_<comision_id>_<cohorte_id>` (ver docentes/views.py). El
        # caché es un LocMemCache de proceso: no lo resetea el rollback
        # transaccional de cada test, así que sobrevive entre tests. Como
        # sqlite reasigna PKs libremente tras cada rollback, otro test que
        # haya corrido antes puede haber dejado una entrada para el mismo par
        # (comision_id, cohorte_id) que este test reutiliza por coincidencia
        # de PKs — y esta vista serviría esos datos viejos en vez de
        # recalcular. En producción los IDs de Comision jamás se reciclan, así
        # que esto no puede pasar ahí: es un artefacto de la suite de tests.
        from django.core.cache import cache
        cache.clear()

        self.docente = _u('doc-prog', es_docente=True)
        self.docente_otro = _u('doc-prog-otro', es_docente=True)
        self.admin = _u('admin-prog', is_staff=True, is_superuser=True)

        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Tarde')
        self.comision.docentes.add(self.docente)

        self.estudiante = _u('alu-prog')
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte)

        self.practica = Practica.objects.create(titulo='Práctica lógica', creada_por=self.docente)
        self.pc = PracticaComision.objects.create(practica=self.practica, comision=self.comision, orden=1)
        ejercicio_1 = Ejercicio.objects.create(enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente)
        ejercicio_2 = Ejercicio.objects.create(enunciado='E2', formula_solucion='q', tipo='formalizacion', creado_por=self.docente)
        self.ejercicio_practica_1 = EjercicioPractica.objects.create(practica=self.practica, ejercicio=ejercicio_1, orden=1)
        EjercicioPractica.objects.create(practica=self.practica, ejercicio=ejercicio_2, orden=2)

    def test_restringe_acceso_a_docente_asignado_o_admin(self):
        self.client.login(username='doc-prog-otro', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))
        self.assertEqual(response.status_code, 403)

        self.client.login(username='admin-prog', password='clave123')
        response_admin = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))
        self.assertEqual(response_admin.status_code, 200)

    def test_muestra_metricas_de_progreso_por_estudiante(self):
        Progreso.objects.create(
            estudiante=self.estudiante,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            ejercicio_practica_actual=self.ejercicio_practica_1,
        )
        # El avance docente se basa en aprobado_docente=True; con ep1 aprobado: 1/2 = 50%
        Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ejercicio_practica_1,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='p',
            es_correcto=True,
            aprobado_docente=True,
        )

        self.client.login(username='doc-prog', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '0/1')      # practicas_completadas/practicas_totales
        self.assertContains(response, 'Ej. 1')    # ejercicio_actual.orden (template: "Ej. N")
        self.assertContains(response, '50,0%')
        self.assertContains(response, reverse('docentes:estudiante_detail', args=[self.estudiante.id]))

    def test_muestra_resuelto_aunque_el_docente_no_haya_revisado(self):
        """El caso de la comisión masiva: nadie corrige a mano y el panel igual informa."""
        Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ejercicio_practica_1,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='p',
            es_correcto=True,
            aprobado_docente=None,
        )

        self.client.login(username='doc-prog', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Resuelto')
        self.assertContains(response, 'Revisado')
        self.assertContains(response, '1/2')  # resueltos_count/total_ejercicios

        # Los porcentajes se verifican en el contexto y no en el HTML: '0,0%' es
        # substring de '50,0%', así que un assertContains pasaría por accidente.
        fila = response.context['progreso_estudiantes'][0]
        self.assertEqual(fila['porcentaje_resuelto'], 50.0)
        self.assertEqual(fila['porcentaje_global'], 0.0)

    def test_analiticas_del_dashboard_se_calculan_sin_error(self):
        """Regresión: comision_detail llama a ocho helpers importados de
        analiticas.views (_ejercicios_mas_dificiles, _estudiantes_en_riesgo,
        _distribucion_intentos, _evolucion_temporal, _errores_sistematicos,
        _silencio_temprano, _concentracion_practica, _velocidad_arranque) con
        el parámetro cohorte_id=<int|None>. Esas funciones fueron renombradas
        a cohorte_ids=<list|None> en analiticas/views.py; si docentes/views.py
        no se actualiza junto con ellas, esta vista tira TypeError y 500."""
        Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ejercicio_practica_1,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='p',
            es_correcto=True,
            aprobado_docente=True,
        )

        self.client.login(username='doc-prog', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        # El dict `analiticas` de la vista se desparrama con **analiticas
        # directo en el contexto del template (no vive bajo una clave
        # 'analiticas'), así que cada resultado es una clave de contexto
        # de primer nivel.
        for clave in (
            'ejercicios_dificiles', 'en_riesgo', 'distribucion_intentos',
            'evolucion_temporal', 'errores_sistematicos', 'silencio',
            'silencio_resumen', 'concentracion', 'velocidad',
        ):
            self.assertIn(clave, response.context)

        # Dato concreto: el único intento (resuelto al primer try) debe
        # aparecer computado en la distribución de intentos del ejercicio 1.
        fila = response.context['distribucion_intentos'][0]
        self.assertEqual(fila['resolvieron'], 1)
        self.assertEqual(fila['promedio'], 1.0)

    def test_resuelto_ignora_intentos_rechazados_por_el_docente(self):
        Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ejercicio_practica_1,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='p',
            es_correcto=True,
            aprobado_docente=False,
        )

        self.client.login(username='doc-prog', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '0/2')


class ProgresoResueltoRecursanteTests(TestCase):
    """Code review (Important I1): _aprobados_por_estudiante_ep y
    _resueltos_por_estudiante_ep indexaban por (estudiante, comision, ep) sin
    cohorte. Un recursante con intentos SOLO en su cohorte vieja, mirado
    desde la cohorte actual (donde Progreso dice que recién arranca), veía
    resueltos_count "prestado" de la camada anterior: la fila del panel
    quedaba contradictoria consigo misma (Progreso: recién empieza / conteo
    de resueltos: ya resolvió)."""

    def setUp(self):
        self.docente = _u('doc-prog-recursa', es_docente=True)
        self.estudiante = _u('est-prog-recursa')

        self.cohorte_vieja = Cohorte.objects.create(anio=2020, cuatrimestre=1)
        self.cohorte_actual = Cohorte.objects.get(anio=2026, cuatrimestre=1)

        self.comision = Comision.objects.create(nombre='IPC Prog Recursa')
        self.comision.docentes.add(self.docente)

        # Vieja primero, actual después: la cohorte "actual" del estudiante en
        # esta comisión (según cohorte_de/inscripción más reciente) es la
        # nueva, la que _resolver_cohorte va a mostrar por default (activa).
        Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte_vieja,
        )
        Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte_actual,
        )

        self.practica = Practica.objects.create(titulo='P-prog-recursa', creada_por=self.docente)
        self.pc = PracticaComision.objects.create(practica=self.practica, comision=self.comision, orden=1)
        ejercicio = Ejercicio.objects.create(
            enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        self.ep = EjercicioPractica.objects.create(practica=self.practica, ejercicio=ejercicio, orden=1)

        # Resuelto Y aprobado SOLO en la camada vieja.
        Intento.objects.create(
            estudiante=self.estudiante, ejercicio_practica=self.ep, practica_comision=self.pc,
            cohorte=self.cohorte_vieja, respuesta_raw='p', es_correcto=True, aprobado_docente=True,
        )

        self.client.login(username='doc-prog-recursa', password='clave123')

    def test_resueltos_count_no_hereda_intentos_de_la_camada_anterior(self):
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        fila = next(
            f for f in response.context['progreso_estudiantes']
            if f['estudiante'].id == self.estudiante.id
        )
        detalle = fila['practicas_detalle'][0]
        # En la cohorte actual (activa, sin ?cohorte= explícito) no hay
        # intentos: resueltos_count debe ser 0, no 1 heredado de la vieja.
        self.assertEqual(detalle['resueltos_count'], 0)
        self.assertFalse(detalle['completada'])
        self.assertEqual(fila['porcentaje_resuelto'], 0.0)
        self.assertEqual(fila['porcentaje_global'], 0.0)

    def test_resueltos_count_ve_la_camada_vieja_si_se_pide_esa_cohorte(self):
        """Contraparte: viendo explícitamente la cohorte vieja, sí debe contar."""
        url = reverse('docentes:comision_detail', args=[self.comision.id])
        response = self.client.get(f'{url}?cohorte={self.cohorte_vieja.pk}')

        self.assertEqual(response.status_code, 200)
        fila = next(
            f for f in response.context['progreso_estudiantes']
            if f['estudiante'].id == self.estudiante.id
        )
        detalle = fila['practicas_detalle'][0]
        self.assertEqual(detalle['resueltos_count'], 1)
        self.assertTrue(detalle['completada'])


class EstudianteDetailIntentosTests(TestCase):
    def setUp(self):
        self.docente = _u('doc-det', es_docente=True)
        self.docente_otro = _u('doc-det-otro', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Mañana')
        self.comision.docentes.add(self.docente)

        self.estudiante = _u('alu-det')
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte)

        self.practica = Practica.objects.create(titulo='Práctica A', creada_por=self.docente)
        self.pc = PracticaComision.objects.create(practica=self.practica, comision=self.comision, orden=1)
        ejercicio = Ejercicio.objects.create(enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente)
        self.ep = EjercicioPractica.objects.create(practica=self.practica, ejercicio=ejercicio, orden=1)

    def test_muestra_ultimo_intento_y_historial_agrupado(self):
        Intento.objects.create(estudiante=self.estudiante, ejercicio_practica=self.ep, practica_comision=self.pc, cohorte=self.cohorte, respuesta_raw='p', es_correcto=True)
        Intento.objects.create(estudiante=self.estudiante, ejercicio_practica=self.ep, practica_comision=self.pc, cohorte=self.cohorte, respuesta_raw='q', es_correcto=False)

        self.client.login(username='doc-det', password='clave123')
        response = self.client.get(reverse('docentes:estudiante_detail', args=[self.estudiante.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'q')
        self.assertContains(response, 'Ver historial (1)')

    def test_restringe_acceso_si_no_docente_de_comision(self):
        self.client.login(username='doc-det-otro', password='clave123')
        response = self.client.get(reverse('docentes:estudiante_detail', args=[self.estudiante.id]))
        self.assertEqual(response.status_code, 403)

    def test_aplica_filtros_por_comision_y_practica(self):
        otra_comision = Comision.objects.create(nombre='IPC Noche')
        otra_comision.docentes.add(self.docente)
        Inscripcion.objects.create(estudiante=self.estudiante, comision=otra_comision, cohorte=self.cohorte)
        otra_practica = Practica.objects.create(titulo='Práctica B', creada_por=self.docente)
        otra_pc = PracticaComision.objects.create(practica=otra_practica, comision=otra_comision, orden=1)
        ejercicio_b = Ejercicio.objects.create(enunciado='E2', formula_solucion='q', tipo='formalizacion', creado_por=self.docente)
        ep_b = EjercicioPractica.objects.create(practica=otra_practica, ejercicio=ejercicio_b, orden=1)
        Intento.objects.create(estudiante=self.estudiante, ejercicio_practica=self.ep, practica_comision=self.pc, cohorte=self.cohorte, respuesta_raw='p', es_correcto=True)
        Intento.objects.create(estudiante=self.estudiante, ejercicio_practica=ep_b, practica_comision=otra_pc, cohorte=self.cohorte, respuesta_raw='q', es_correcto=True)

        self.client.login(username='doc-det', password='clave123')
        response = self.client.get(
            reverse('docentes:estudiante_detail', args=[self.estudiante.id]),
            {'comision': otra_comision.id, 'pc': otra_pc.id},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Práctica B')
        self.assertNotContains(response, 'Práctica A')


class EjercicioPracticaOrdenamientoTests(TestCase):
    def setUp(self):
        self.docente = _u('doc-orden', es_docente=True)
        self.comision = Comision.objects.create(nombre='IPC Orden')
        self.comision.docentes.add(self.docente)

        self.practica = Practica.objects.create(titulo='Práctica Orden', creada_por=self.docente)
        PracticaComision.objects.create(practica=self.practica, comision=self.comision, orden=1)
        self.ej1 = Ejercicio.objects.create(enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente)
        self.ej2 = Ejercicio.objects.create(enunciado='E2', formula_solucion='q', tipo='formalizacion', creado_por=self.docente)
        self.ej3 = Ejercicio.objects.create(enunciado='E3', formula_solucion='r', tipo='formalizacion', creado_por=self.docente)

        self.ep1 = EjercicioPractica.objects.create(practica=self.practica, ejercicio=self.ej1, orden=1)
        self.ep2 = EjercicioPractica.objects.create(practica=self.practica, ejercicio=self.ej2, orden=2)

        self.client.login(username='doc-orden', password='clave123')

    def test_agregar_en_orden_existente_desplaza_siguientes(self):
        response = self.client.post(
            reverse('docentes:ejercicio_practica_add', args=[self.practica.id]),
            data={'practica': self.practica.id, 'ejercicio': self.ej3.id, 'orden': 2},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        ordenes = list(
            EjercicioPractica.objects.filter(practica=self.practica)
            .order_by('orden')
            .values_list('ejercicio_id', 'orden')
        )
        self.assertEqual(ordenes, [(self.ej1.id, 1), (self.ej3.id, 2), (self.ej2.id, 3)])

    def test_eliminar_en_medio_reordena_los_siguientes(self):
        ep3 = EjercicioPractica.objects.create(practica=self.practica, ejercicio=self.ej3, orden=3)

        response = self.client.post(
            reverse('docentes:ejercicio_practica_remove', args=[self.ep2.id]),
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(EjercicioPractica.objects.filter(pk=self.ep2.id).exists())
        ep3.refresh_from_db()
        self.assertEqual(ep3.orden, 2)

    def test_reordenar_con_flechas_no_da_error_y_actualiza_orden(self):
        response = self.client.post(
            reverse('docentes:ejercicio_practica_reorder', args=[self.ep2.id]),
            data={'direccion': 'subir'},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.ep1.refresh_from_db()
        self.ep2.refresh_from_db()
        self.assertEqual(self.ep2.orden, 1)
        self.assertEqual(self.ep1.orden, 2)


class EjercicioFormVolverTests(TestCase):
    def setUp(self):
        self.docente = _u('doc-volver', es_docente=True)
        self.comision = Comision.objects.create(nombre='IPC Volver')
        self.comision.docentes.add(self.docente)
        self.practica = Practica.objects.create(titulo='Práctica Volver', creada_por=self.docente)
        PracticaComision.objects.create(practica=self.practica, comision=self.comision, orden=1)
        self.ejercicio = Ejercicio.objects.create(
            enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente
        )
        EjercicioPractica.objects.create(practica=self.practica, ejercicio=self.ejercicio, orden=1)
        self.client.login(username='doc-volver', password='clave123')

    def test_practica_detail_envia_next_en_links_a_crear_y_editar(self):
        response = self.client.get(reverse('docentes:practica_detail', args=[self.practica.id]))

        self.assertEqual(response.status_code, 200)
        next_url = reverse('docentes:practica_detail', args=[self.practica.id])
        self.assertContains(
            response,
            f"{reverse('docentes:ejercicio_create')}?practica={self.practica.id}&next={next_url}",
        )
        self.assertContains(
            response,
            f"{reverse('docentes:ejercicio_edit', args=[self.ejercicio.id])}?next={next_url}",
        )

    def test_formulario_usa_next_para_boton_volver(self):
        next_url = reverse('docentes:practica_detail', args=[self.practica.id])

        response = self.client.get(reverse('docentes:ejercicio_create'), {'next': next_url})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'<input type="hidden" name="next" value="{next_url}">', html=True)
        self.assertContains(response, f'<a class="btn btn-secondary" href="{next_url}">Volver</a>', html=True)


class PracticaDetailPermisosTests(TestCase):
    """Criterio de acceso a practica_detail.

    Lectura: staff, autor, práctica pública, o docente de una comisión donde
    esté asignada. Escritura: staff, docente de una comisión donde esté
    asignada, o el autor mientras ninguna comisión la use. Ser pública NO
    habilita escritura.
    """

    def setUp(self):
        self.autor = _u('doc-autor', es_docente=True)
        self.otro = _u('doc-otro', es_docente=True)
        self.colega = _u('doc-colega', es_docente=True)
        self.admin = _u('admin-practicas', es_docente=True, is_staff=True)

        self.comision = Comision.objects.create(nombre='IPC Permisos')
        self.comision.docentes.add(self.autor, self.colega)

        # Práctica del autor sin ninguna comisión asignada
        self.huerfana = Practica.objects.create(titulo='Huérfana', creada_por=self.autor)

        # Práctica pública del autor, tampoco asignada a comisiones de otros
        self.publica = Practica.objects.create(
            titulo='Pública', creada_por=self.autor, es_publica=True
        )

        # Práctica privada del autor, asignada a una comisión donde enseña colega
        self.asignada = Practica.objects.create(titulo='Asignada', creada_por=self.autor)
        PracticaComision.objects.create(
            practica=self.asignada, comision=self.comision, orden=1
        )

        self.ejercicio = Ejercicio.objects.create(
            enunciado='E público', formula_solucion='p', tipo='formalizacion',
            creado_por=self.autor, es_publico=True,
        )
        self.ep_publica = EjercicioPractica.objects.create(
            practica=self.publica, ejercicio=self.ejercicio, orden=1
        )

    def _url(self, practica):
        return reverse('docentes:practica_detail', args=[practica.id])

    # ── Lectura ──────────────────────────────────────────────────────
    def test_autor_ve_su_practica_sin_comision_asignada(self):
        self.client.login(username='doc-autor', password='clave123')

        response = self.client.get(self._url(self.huerfana))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['puede_editar'])

    def test_docente_ajeno_ve_practica_publica_en_solo_lectura(self):
        self.client.login(username='doc-otro', password='clave123')

        response = self.client.get(self._url(self.publica))

        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['puede_editar'])

    def test_docente_ajeno_no_ve_practica_privada_de_otro(self):
        self.client.login(username='doc-otro', password='clave123')

        response = self.client.get(self._url(self.asignada))

        self.assertEqual(response.status_code, 403)

    def test_docente_de_la_comision_ve_y_edita(self):
        self.client.login(username='doc-colega', password='clave123')

        response = self.client.get(self._url(self.asignada))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['puede_editar'])

    def test_staff_ve_cualquier_practica(self):
        self.client.login(username='admin-practicas', password='clave123')

        response = self.client.get(self._url(self.huerfana))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['puede_editar'])

    # ── Acciones ocultas en modo lectura ─────────────────────────────
    def test_lectura_no_muestra_acciones_de_escritura(self):
        self.client.login(username='doc-otro', password='clave123')

        response = self.client.get(self._url(self.publica))

        self.assertNotContains(
            response, reverse('docentes:ejercicio_practica_add', args=[self.publica.id])
        )
        self.assertNotContains(
            response, reverse('docentes:ejercicio_practica_remove', args=[self.ep_publica.id])
        )
        self.assertNotContains(
            response, reverse('docentes:ejercicio_practica_reorder', args=[self.ep_publica.id])
        )
        self.assertNotContains(
            response, reverse('docentes:ejercicio_edit', args=[self.ejercicio.id])
        )

    def test_autor_si_ve_acciones_de_escritura(self):
        self.client.login(username='doc-autor', password='clave123')

        response = self.client.get(self._url(self.publica))

        self.assertContains(
            response, reverse('docentes:ejercicio_practica_add', args=[self.publica.id])
        )
        self.assertContains(
            response, reverse('docentes:ejercicio_practica_reorder', args=[self.ep_publica.id])
        )

    # ── Escritura ────────────────────────────────────────────────────
    def test_lectura_publica_no_habilita_agregar_ejercicio(self):
        self.client.login(username='doc-otro', password='clave123')
        ejercicio_propio = Ejercicio.objects.create(
            enunciado='E otro', formula_solucion='q', tipo='formalizacion',
            creado_por=self.otro,
        )

        response = self.client.post(
            reverse('docentes:ejercicio_practica_add', args=[self.publica.id]),
            data={'ejercicio': ejercicio_propio.id, 'orden': 2},
        )

        self.assertEqual(response.status_code, 403)
        self.assertEqual(EjercicioPractica.objects.filter(practica=self.publica).count(), 1)

    def test_lectura_publica_no_habilita_quitar_ni_reordenar(self):
        self.client.login(username='doc-otro', password='clave123')

        remove = self.client.post(
            reverse('docentes:ejercicio_practica_remove', args=[self.ep_publica.id])
        )
        reorder = self.client.post(
            reverse('docentes:ejercicio_practica_reorder', args=[self.ep_publica.id]),
            data={'direccion': 'subir'},
        )

        self.assertEqual(remove.status_code, 403)
        self.assertEqual(reorder.status_code, 403)
        self.assertTrue(EjercicioPractica.objects.filter(pk=self.ep_publica.id).exists())

    def test_autor_puede_agregar_ejercicio_a_practica_sin_comision(self):
        self.client.login(username='doc-autor', password='clave123')

        response = self.client.post(
            reverse('docentes:ejercicio_practica_add', args=[self.huerfana.id]),
            data={'ejercicio': self.ejercicio.id, 'orden': 1},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(
            EjercicioPractica.objects.filter(
                practica=self.huerfana, ejercicio=self.ejercicio
            ).exists()
        )

    # ── El autor pierde la escritura si la usa una comisión ajena ────
    def test_autor_no_edita_practica_que_usa_una_comision_ajena(self):
        """La Practica es canónica: importarla la comparte, no la copia.

        Si el autor no enseña en ninguna de las comisiones que la usan, no
        puede tocar sus ejercicios: quitar uno borra en cascada los Intento
        de los estudiantes de esa comisión.
        """
        comision_ajena = Comision.objects.create(nombre='IPC Ajena')
        comision_ajena.docentes.add(self.otro)
        PracticaComision.objects.create(
            practica=self.publica, comision=comision_ajena, orden=1
        )

        self.client.login(username='doc-autor', password='clave123')
        detail = self.client.get(self._url(self.publica))
        remove = self.client.post(
            reverse('docentes:ejercicio_practica_remove', args=[self.ep_publica.id])
        )

        self.assertEqual(detail.status_code, 200)
        self.assertFalse(detail.context['puede_editar'])
        self.assertEqual(remove.status_code, 403)
        self.assertTrue(EjercicioPractica.objects.filter(pk=self.ep_publica.id).exists())

    def test_autor_sigue_leyendo_su_practica_privada_en_comision_ajena(self):
        comision_ajena = Comision.objects.create(nombre='IPC Ajena')
        comision_ajena.docentes.add(self.otro)
        PracticaComision.objects.create(
            practica=self.asignada, comision=comision_ajena, orden=2
        )
        self.comision.docentes.remove(self.autor)

        self.client.login(username='doc-autor', password='clave123')
        response = self.client.get(self._url(self.asignada))

        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['puede_editar'])

    def test_autor_que_ensena_en_la_comision_conserva_la_escritura(self):
        """Caso normal: el autor enseña donde asignó la práctica."""
        self.client.login(username='doc-autor', password='clave123')

        response = self.client.get(self._url(self.asignada))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['puede_editar'])


class PracticaOrigenTests(TestCase):
    """Cubre practica_import (crea PracticaComision) y _desanclar_si_derivada.

    Con el modelo PracticaComision, importar una práctica ya no crea una copia:
    solo crea un vínculo PracticaComision entre la Practica canónica y la
    Comision destino.  El mecanismo practica_origen sigue vigente para prácticas
    que fueron copiadas manualmente y luego editadas.
    """

    def setUp(self):
        self.docente = _u('doc-origen', es_docente=True)
        self.comision_a = Comision.objects.create(nombre='IPC A')
        self.comision_b = Comision.objects.create(nombre='IPC B')
        self.comision_a.docentes.add(self.docente)
        self.comision_b.docentes.add(self.docente)

        # Práctica canónica con dos ejercicios, asignada a comision_a
        self.practica_orig = Practica.objects.create(
            titulo='Práctica Original',
            creada_por=self.docente,
            es_publica=True,
        )
        self.pc_a = PracticaComision.objects.create(
            practica=self.practica_orig, comision=self.comision_a, orden=1
        )
        self.ej1 = Ejercicio.objects.create(enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente)
        self.ej2 = Ejercicio.objects.create(enunciado='E2', formula_solucion='q', tipo='formalizacion', creado_por=self.docente)
        EjercicioPractica.objects.create(practica=self.practica_orig, ejercicio=self.ej1, orden=1)
        EjercicioPractica.objects.create(practica=self.practica_orig, ejercicio=self.ej2, orden=2)

        self.client.login(username='doc-origen', password='clave123')

    def _importar(self, comision=None, orden=1):
        """Crea un PracticaComision via la vista practica_import."""
        comision = comision or self.comision_b
        return self.client.post(
            reverse('docentes:practica_import', args=[comision.id]),
            data={'practica_origen': self.practica_orig.id, 'orden_destino': orden},
            follow=True,
        )

    def _crear_derivada(self, comision=None, orden=1):
        """Crea manualmente una Practica con practica_origen seteado + su PC."""
        comision = comision or self.comision_b
        derivada = Practica.objects.create(
            titulo=self.practica_orig.titulo,
            creada_por=self.docente,
            practica_origen=self.practica_orig,
        )
        EjercicioPractica.objects.create(practica=derivada, ejercicio=self.ej1, orden=1)
        EjercicioPractica.objects.create(practica=derivada, ejercicio=self.ej2, orden=2)
        pc = PracticaComision.objects.create(practica=derivada, comision=comision, orden=orden)
        return derivada, pc

    def test_importar_crea_practica_comision(self):
        """practica_import crea un PracticaComision, no una copia de Practica."""
        self._importar()

        self.assertTrue(
            PracticaComision.objects.filter(
                practica=self.practica_orig, comision=self.comision_b
            ).exists()
        )
        # No se creó una nueva Practica
        self.assertEqual(Practica.objects.count(), 1)

    def test_practica_importada_tiene_ejercicios_en_orden(self):
        """Tras importar, la práctica canónica mantiene sus ejercicios en orden."""
        self._importar()

        eps = list(
            self.practica_orig.ejercicio_practicas
            .order_by('orden')
            .values_list('ejercicio_id', 'orden')
        )
        self.assertEqual(eps, [(self.ej1.id, 1), (self.ej2.id, 2)])

    def test_banco_excluye_practicas_derivadas(self):
        """El banco filtra con practica_origen__isnull=True: excluye derivadas."""
        derivada, _ = self._crear_derivada()
        self.assertIsNotNone(derivada.practica_origen_id)

        banco = Practica.objects.filter(practica_origen__isnull=True)
        self.assertIn(self.practica_orig, banco)
        self.assertNotIn(derivada, banco)

    def test_editar_practica_con_origen_la_desancla(self):
        """Editar (vía practica_edit) una práctica derivada limpia practica_origen."""
        derivada, pc = self._crear_derivada()
        self.assertIsNotNone(derivada.practica_origen_id)

        self.client.post(
            reverse('docentes:practica_edit', args=[pc.id]),  # pc_id
            data={
                'titulo': 'Título modificado',
                'es_publica': False,
                'orden': 1,
                'comision': self.comision_b.id,
            },
            follow=True,
        )

        derivada.refresh_from_db()
        self.assertIsNone(derivada.practica_origen_id)

    def test_agregar_ejercicio_desancla_practica(self):
        """Agregar un ejercicio a una práctica derivada limpia practica_origen."""
        derivada, _ = self._crear_derivada()
        ej3 = Ejercicio.objects.create(enunciado='E3', formula_solucion='r', tipo='formalizacion', creado_por=self.docente)

        self.client.post(
            reverse('docentes:ejercicio_practica_add', args=[derivada.id]),
            data={'practica': derivada.id, 'ejercicio': ej3.id, 'orden': 3},
            follow=True,
        )

        derivada.refresh_from_db()
        self.assertIsNone(derivada.practica_origen_id)

    def test_eliminar_ejercicio_desancla_practica(self):
        """Quitar un ejercicio de una práctica derivada limpia practica_origen."""
        derivada, _ = self._crear_derivada()
        ep = derivada.ejercicio_practicas.order_by('orden').first()

        self.client.post(
            reverse('docentes:ejercicio_practica_remove', args=[ep.id]),
            follow=True,
        )

        derivada.refresh_from_db()
        self.assertIsNone(derivada.practica_origen_id)

    def test_reordenar_ejercicio_desancla_practica(self):
        """Reordenar ejercicios en una práctica derivada limpia practica_origen."""
        derivada, _ = self._crear_derivada()
        ep2 = derivada.ejercicio_practicas.order_by('orden').last()

        self.client.post(
            reverse('docentes:ejercicio_practica_reorder', args=[ep2.id]),
            data={'direccion': 'subir'},
            follow=True,
        )

        derivada.refresh_from_db()
        self.assertIsNone(derivada.practica_origen_id)


class IntentoAprobacionUpdateTests(TestCase):
    def setUp(self):
        self.docente = _u('doc-aprueba', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Tarde')
        self.comision.docentes.add(self.docente)

        self.estudiante = _u('alu-aprueba')
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte)

        self.practica = Practica.objects.create(titulo='Práctica Corrección')
        self.pc = PracticaComision.objects.create(practica=self.practica, comision=self.comision, orden=1)
        ejercicio = Ejercicio.objects.create(
            enunciado='E1',
            formula_solucion='p',
            tipo='formalizacion',
            creado_por=self.docente,
        )
        self.ep = EjercicioPractica.objects.create(practica=self.practica, ejercicio=ejercicio, orden=1)
        self.intento = Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ep,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='p',
            es_correcto=True,
            aprobado_docente=None,
            comentario_docente='',
        )

        self.client.login(username='doc-aprueba', password='clave123')

    def test_rechazar_sin_comentario_desde_correccion_pendiente_muestra_error_y_no_actualiza(self):
        next_url = reverse('docentes:correccion_pendiente')

        response = self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[self.intento.id]),
            data={
                'accion': 'rechazar',
                'comentario_docente': '   ',
                'next': next_url,
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            'Para rechazar un intento es necesario escribir una devolución. '
            'La devolución argumentada es parte del acompañamiento pedagógico.',
        )
        self.intento.refresh_from_db()
        self.assertIsNone(self.intento.aprobado_docente)
        self.assertEqual(self.intento.comentario_docente, '')

    def test_mensaje_error_por_rechazo_sin_comentario_es_consistente(self):
        response = self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[self.intento.id]),
            data={
                'accion': 'rechazar',
                'comentario_docente': '',
                'next': reverse('docentes:estudiante_detail', args=[self.estudiante.id]),
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            'Para rechazar un intento es necesario escribir una devolución. '
            'La devolución argumentada es parte del acompañamiento pedagógico.',
        )


class IntentoAprobacionDesbloqueoTests(TestCase):
    """Cuando el docente aprueba un intento incorrecto, debe desbloquearse el siguiente ejercicio."""

    def setUp(self):
        self.docente = _u('doc-desbloqueo', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Mañana')
        self.comision.docentes.add(self.docente)

        self.estudiante = _u('alu-desbloqueo')
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte)

        self.practica = Practica.objects.create(titulo='Práctica Desbloqueo')
        self.pc = PracticaComision.objects.create(practica=self.practica, comision=self.comision, orden=1)

        ej1 = Ejercicio.objects.create(enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente)
        ej2 = Ejercicio.objects.create(enunciado='E2', formula_solucion='q', tipo='formalizacion', creado_por=self.docente)
        self.ep1 = EjercicioPractica.objects.create(practica=self.practica, ejercicio=ej1, orden=1)
        self.ep2 = EjercicioPractica.objects.create(practica=self.practica, ejercicio=ej2, orden=2)

        self.intento_incorrecto = Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ep1,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='q->p',
            es_correcto=False,
            aprobado_docente=None,
        )
        self.progreso = Progreso.objects.create(
            estudiante=self.estudiante,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            ejercicio_practica_actual=self.ep1,
        )

        self.client.login(username='doc-desbloqueo', password='clave123')

    def test_aprobar_intento_incorrecto_avanza_progreso(self):
        self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[self.intento_incorrecto.id]),
            data={'accion': 'aprobar'},
        )

        self.progreso.refresh_from_db()
        self.assertEqual(self.progreso.ejercicio_practica_actual, self.ep2)

    def test_aprobar_intento_incorrecto_ultimo_ejercicio_completa_practica(self):
        self.progreso.ejercicio_practica_actual = self.ep2
        self.progreso.save()

        intento_ep2 = Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ep2,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='p->q',
            es_correcto=False,
            aprobado_docente=None,
        )

        self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[intento_ep2.id]),
            data={'accion': 'aprobar'},
        )

        self.progreso.refresh_from_db()
        self.assertIsNone(self.progreso.ejercicio_practica_actual)

    def test_aprobar_intento_correcto_no_duplica_avance(self):
        intento_correcto = Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ep1,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='p',
            es_correcto=True,
            aprobado_docente=None,
        )
        self.progreso.ejercicio_practica_actual = self.ep2
        self.progreso.save()

        self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[intento_correcto.id]),
            data={'accion': 'aprobar'},
        )

        self.progreso.refresh_from_db()
        self.assertEqual(self.progreso.ejercicio_practica_actual, self.ep2)

    def test_rechazar_intento_incorrecto_no_avanza_progreso(self):
        self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[self.intento_incorrecto.id]),
            data={'accion': 'rechazar', 'comentario_docente': 'Fórmula incorrecta.'},
        )

        self.progreso.refresh_from_db()
        self.assertEqual(self.progreso.ejercicio_practica_actual, self.ep1)


class ParcialTests(TestCase):
    def setUp(self):
        self.docente = _u('doc_parcial', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Mañana')
        self.comision.docentes.add(self.docente)
        self.est1 = _u('est1_parcial')
        self.est2 = _u('est2_parcial')
        Inscripcion.objects.create(estudiante=self.est1, comision=self.comision, cohorte=self.cohorte)
        Inscripcion.objects.create(estudiante=self.est2, comision=self.comision, cohorte=self.cohorte)

    def test_crear_parcial_precrea_notas_para_todos_los_inscriptos(self):
        self.client.login(username='doc_parcial', password='clave123')
        resp = self.client.post(
            reverse('docentes:parcial_create', kwargs={'comision_id': self.comision.id}),
            {'nombre': 'Parcial 1', 'fecha': '2026-06-15', 'puntaje_total': '10.00'},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Parcial.objects.count(), 1)
        parcial = Parcial.objects.get()
        self.assertEqual(parcial.comision, self.comision)
        self.assertEqual(NotaParcial.objects.filter(parcial=parcial).count(), 2)
        notas = NotaParcial.objects.filter(parcial=parcial)
        self.assertTrue(all(n.puntaje is None for n in notas))

    def test_docente_externo_no_puede_crear_parcial(self):
        otro = _u('otro_docente', es_docente=True)
        self.client.login(username='otro_docente', password='clave123')
        resp = self.client.post(
            reverse('docentes:parcial_create', kwargs={'comision_id': self.comision.id}),
            {'nombre': 'Parcial 1', 'fecha': '2026-06-15', 'puntaje_total': '10.00'},
        )
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Parcial.objects.count(), 0)

    def test_crear_parcial_sin_cohorte_activa_no_crea_nada(self):
        """Sin cohorte activa, `_require_cohorte_activa` corta con 403 en vez
        de reventar con IntegrityError o crear el parcial en una camada
        cerrada."""
        Cohorte.objects.filter(activa=True).update(activa=False)
        self.client.login(username='doc_parcial', password='clave123')

        resp = self.client.post(
            reverse('docentes:parcial_create', kwargs={'comision_id': self.comision.id}),
            {'nombre': 'Parcial 1', 'fecha': '2026-06-15', 'puntaje_total': '10.00'},
        )

        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Parcial.objects.count(), 0)

    def test_guardar_notas_parciales(self):
        self.client.login(username='doc_parcial', password='clave123')
        parcial = Parcial.objects.create(
            comision=self.comision, cohorte=self.cohorte, nombre='P1',
            fecha=datetime.date(2026, 6, 15), puntaje_total='10.00',
        )
        nota1 = NotaParcial.objects.create(parcial=parcial, estudiante=self.est1)
        nota2 = NotaParcial.objects.create(parcial=parcial, estudiante=self.est2)

        resp = self.client.post(
            reverse('docentes:parcial_notas', kwargs={'parcial_id': parcial.id}),
            {
                'form-TOTAL_FORMS': '2',
                'form-INITIAL_FORMS': '2',
                'form-MIN_NUM_FORMS': '0',
                'form-MAX_NUM_FORMS': '1000',
                f'form-0-id': str(nota1.id),
                f'form-0-puntaje': '7.50',
                f'form-1-id': str(nota2.id),
                f'form-1-puntaje': '',
            },
        )
        self.assertEqual(resp.status_code, 302)
        nota1.refresh_from_db()
        nota2.refresh_from_db()
        self.assertEqual(nota1.puntaje, Decimal('7.50'))
        self.assertIsNone(nota2.puntaje)

    def test_eliminar_parcial_elimina_notas(self):
        self.client.login(username='doc_parcial', password='clave123')
        parcial = Parcial.objects.create(
            comision=self.comision, cohorte=self.cohorte, nombre='P1',
            fecha=datetime.date(2026, 6, 15), puntaje_total='10.00',
        )
        NotaParcial.objects.create(parcial=parcial, estudiante=self.est1)
        resp = self.client.post(
            reverse('docentes:parcial_delete', kwargs={'parcial_id': parcial.id}),
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Parcial.objects.count(), 0)
        self.assertEqual(NotaParcial.objects.count(), 0)

    def test_notaparcial_tiene_campo_ausente_y_nota_logica_parcial(self):
        parcial = Parcial.objects.create(
            comision=self.comision, cohorte=self.cohorte, nombre='P1',
            fecha=datetime.date(2026, 6, 15), puntaje_total='10.00',
        )
        nota = NotaParcial.objects.create(parcial=parcial, estudiante=self.est1)
        self.assertFalse(nota.ausente)               # default False
        self.assertIsNone(nota.nota_logica_parcial)  # default null

    def test_parcial_tiene_umbrales_con_defaults(self):
        parcial = Parcial.objects.create(
            comision=self.comision, cohorte=self.cohorte, nombre='P1',
            fecha=datetime.date(2026, 6, 15), puntaje_total='10.00',
        )
        from decimal import Decimal
        self.assertEqual(parcial.umbral_aprobacion, Decimal('4.00'))
        self.assertEqual(parcial.umbral_promocion, Decimal('7.00'))

    def test_ausente_marca_correctamente(self):
        from decimal import Decimal
        from django.core.exceptions import ValidationError
        parcial = Parcial.objects.create(
            comision=self.comision, cohorte=self.cohorte, nombre='P1',
            fecha=datetime.date(2026, 6, 15), puntaje_total='10.00',
        )
        nota = NotaParcial.objects.create(parcial=parcial, estudiante=self.est1)
        nota.ausente = True
        nota.save()
        nota.refresh_from_db()
        self.assertTrue(nota.ausente)
        self.assertIsNone(nota.puntaje)
        self.assertIsNone(nota.nota_logica_parcial)

        # Verify the invariant is enforced: ausente=True con puntaje → falla
        nota2 = NotaParcial(
            parcial=parcial,
            estudiante=self.est2,
            ausente=True,
            puntaje=Decimal('7.50'),
        )
        with self.assertRaises(ValidationError):
            nota2.full_clean()

        # Y también con nota_logica_parcial
        nota3 = NotaParcial(
            parcial=parcial,
            estudiante=self.est2,
            ausente=True,
            nota_logica_parcial=Decimal('4.00'),
        )
        with self.assertRaises(ValidationError):
            nota3.full_clean()

    def test_guardar_ausente_pone_ausente_true_y_notas_null(self):
        self.client.login(username='doc_parcial', password='clave123')
        parcial = Parcial.objects.create(
            comision=self.comision, cohorte=self.cohorte, nombre='P1',
            fecha=datetime.date(2026, 6, 15), puntaje_total='10.00',
        )
        nota1 = NotaParcial.objects.create(parcial=parcial, estudiante=self.est1)
        nota2 = NotaParcial.objects.create(parcial=parcial, estudiante=self.est2)

        resp = self.client.post(
            reverse('docentes:parcial_notas', kwargs={'parcial_id': parcial.id}),
            {
                'form-TOTAL_FORMS': '2',
                'form-INITIAL_FORMS': '2',
                'form-MIN_NUM_FORMS': '0',
                'form-MAX_NUM_FORMS': '1000',
                f'form-0-id': str(nota1.id),
                f'form-0-ausente': 'on',        # checkbox marcado
                f'form-0-puntaje': '',
                f'form-0-nota_logica_parcial': '',
                f'form-1-id': str(nota2.id),
                f'form-1-puntaje': '7.50',
                f'form-1-nota_logica_parcial': '5.00',
            },
        )
        self.assertEqual(resp.status_code, 302)
        nota1.refresh_from_db()
        nota2.refresh_from_db()
        self.assertTrue(nota1.ausente)
        self.assertIsNone(nota1.puntaje)
        self.assertEqual(nota2.puntaje, Decimal('7.50'))
        self.assertEqual(nota2.nota_logica_parcial, Decimal('5.00'))

    def test_conteo_cargadas_incluye_ausentes(self):
        self.client.login(username='doc_parcial', password='clave123')
        parcial = Parcial.objects.create(
            comision=self.comision, cohorte=self.cohorte, nombre='P1',
            fecha=datetime.date(2026, 6, 15), puntaje_total='10.00',
        )
        NotaParcial.objects.create(parcial=parcial, estudiante=self.est1, ausente=True)
        NotaParcial.objects.create(parcial=parcial, estudiante=self.est2)

        resp = self.client.get(
            reverse('docentes:parcial_notas', kwargs={'parcial_id': parcial.id})
        )
        self.assertEqual(resp.status_code, 200)
        # 1 ausente cuenta como cargada, 1 pendiente
        self.assertEqual(resp.context['cargadas'], 1)
        self.assertEqual(resp.context['total'], 2)


class ParcialNotasExportarTests(TestCase):
    def setUp(self):
        self.docente = _u('doc_export', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Tarde')
        self.comision.docentes.add(self.docente)
        self.est1 = _u('est1_export', first_name='Ana', last_name='Aguirre')
        self.est2 = _u('est2_export', first_name='Bruno', last_name='Benítez')
        Inscripcion.objects.create(estudiante=self.est1, comision=self.comision, cohorte=self.cohorte)
        Inscripcion.objects.create(estudiante=self.est2, comision=self.comision, cohorte=self.cohorte)
        EncuestaEstudiante.objects.create(
            estudiante=self.est1, dni='30111222', carrera='abogacia',
        )
        self.parcial = Parcial.objects.create(
            comision=self.comision, cohorte=self.cohorte, nombre='P1',
            fecha=datetime.date(2026, 6, 15), puntaje_total='10.00',
        )
        NotaParcial.objects.create(
            parcial=self.parcial, estudiante=self.est1,
            puntaje=Decimal('7.50'), nota_logica_parcial=Decimal('5.00'),
        )
        NotaParcial.objects.create(
            parcial=self.parcial, estudiante=self.est2, ausente=True,
        )

    def _exportar(self):
        return self.client.get(
            reverse('docentes:parcial_notas_exportar', kwargs={'parcial_id': self.parcial.id})
        )

    def test_exporta_xlsx_con_columnas_y_notas(self):
        self.client.login(username='doc_export', password='clave123')
        resp = self._exportar()
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        worksheet = load_workbook(BytesIO(resp.content)).active
        rows = list(worksheet.iter_rows(values_only=True))
        self.assertEqual(
            rows[0],
            ('DNI', 'Apellido/s', 'Nombre/s', 'Carrera', 'Nota global', 'Nota lógica'),
        )
        # Ordenado por apellido: Aguirre primero, Benítez después
        self.assertEqual(rows[1], ('30111222', 'Aguirre', 'Ana', 'Abogacía', 7.5, 5))
        # Ausente sin encuesta: DNI/carrera vacíos, "Ausente" en ambas notas
        self.assertEqual(rows[2], (None, 'Benítez', 'Bruno', None, 'Ausente', 'Ausente'))

    def test_nota_pendiente_queda_vacia(self):
        NotaParcial.objects.filter(estudiante=self.est2).update(ausente=False)
        self.client.login(username='doc_export', password='clave123')
        worksheet = load_workbook(BytesIO(self._exportar().content)).active
        rows = list(worksheet.iter_rows(values_only=True))
        self.assertEqual(rows[2][4], None)
        self.assertEqual(rows[2][5], None)

    def test_docente_externo_no_puede_exportar(self):
        _u('otro_doc_export', es_docente=True)
        self.client.login(username='otro_doc_export', password='clave123')
        resp = self._exportar()
        self.assertEqual(resp.status_code, 403)


class ResueltosPorEstudianteEpTests(TestCase):
    """El índice de resueltos no depende de la corrección manual del docente."""

    def setUp(self):
        self.docente = _u('doc-resueltos', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Resueltos')
        self.comision.docentes.add(self.docente)
        self.estudiante = _u('alu-resueltos')
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte)

        practica = Practica.objects.create(titulo='P', creada_por=self.docente)
        self.pc = PracticaComision.objects.create(
            practica=practica, comision=self.comision, orden=1,
        )
        ejercicio = Ejercicio.objects.create(
            enunciado='E', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        self.ep = EjercicioPractica.objects.create(
            practica=practica, ejercicio=ejercicio, orden=1,
        )

    def _intento(self, es_correcto, aprobado_docente):
        Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=self.ep,
            practica_comision=self.pc,
            cohorte=self.cohorte,
            respuesta_raw='p',
            es_correcto=es_correcto,
            aprobado_docente=aprobado_docente,
        )

    def _resueltos(self):
        from docentes.views import _resueltos_por_estudiante_ep
        return _resueltos_por_estudiante_ep([self.comision.id], [self.estudiante.id])

    def test_cuenta_correcto_sin_revisar(self):
        from docentes.views import _aprobados_por_estudiante_ep
        self._intento(es_correcto=True, aprobado_docente=None)

        self.assertIn((self.estudiante.id, self.comision.id, self.ep.id), self._resueltos())
        # El índice de aprobados, en cambio, no lo cuenta: ése es el problema
        # que motiva este cambio.
        self.assertNotIn(
            (self.estudiante.id, self.comision.id, self.ep.id),
            _aprobados_por_estudiante_ep([self.comision.id], [self.estudiante.id]),
        )

    def test_no_cuenta_correcto_rechazado(self):
        self._intento(es_correcto=True, aprobado_docente=False)
        self.assertNotIn((self.estudiante.id, self.comision.id, self.ep.id), self._resueltos())

    def test_cuenta_incorrecto_aprobado(self):
        self._intento(es_correcto=False, aprobado_docente=True)
        self.assertIn((self.estudiante.id, self.comision.id, self.ep.id), self._resueltos())

    def test_no_cuenta_incorrecto_sin_revisar(self):
        self._intento(es_correcto=False, aprobado_docente=None)
        self.assertNotIn((self.estudiante.id, self.comision.id, self.ep.id), self._resueltos())

    def test_devuelve_vacio_sin_estudiantes_o_comisiones(self):
        from docentes.views import _resueltos_por_estudiante_ep
        self._intento(es_correcto=True, aprobado_docente=None)

        self.assertEqual(_resueltos_por_estudiante_ep([], [self.estudiante.id]), set())
        self.assertEqual(_resueltos_por_estudiante_ep([self.comision.id], []), set())

class AvanceCruzadoEntreComisionesTests(TestCase):
    """Una misma Practica canónica puede estar en dos comisiones (import)."""

    def setUp(self):
        self.docente = _u('doc-cruce', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision_a = Comision.objects.create(nombre='Comisión A')
        self.comision_b = Comision.objects.create(nombre='Comisión B')
        self.comision_a.docentes.add(self.docente)
        self.comision_b.docentes.add(self.docente)

        # Estudiante inscripto en AMBAS comisiones
        self.estudiante = _u('alu-cruce')
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision_a, cohorte=self.cohorte)
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision_b, cohorte=self.cohorte)

        # La MISMA practica canónica asignada a ambas (lo que hace practica_importar)
        practica = Practica.objects.create(titulo='P compartida', creada_por=self.docente)
        self.pc_a = PracticaComision.objects.create(
            practica=practica, comision=self.comision_a, orden=1,
        )
        self.pc_b = PracticaComision.objects.create(
            practica=practica, comision=self.comision_b, orden=1,
        )
        ejercicio = Ejercicio.objects.create(
            enunciado='E', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        self.ep = EjercicioPractica.objects.create(
            practica=practica, ejercicio=ejercicio, orden=1,
        )

    def test_resuelto_en_A_no_debe_contar_en_B(self):
        from docentes.views import _resueltos_por_estudiante_ep, _calcular_avance_promedio
        # Resuelve SOLO en la comisión A
        Intento.objects.create(
            estudiante=self.estudiante, ejercicio_practica=self.ep,
            practica_comision=self.pc_a, cohorte=self.cohorte,
            respuesta_raw='p', es_correcto=True, aprobado_docente=None,
        )
        ids = [self.comision_a.id, self.comision_b.id]
        resueltos = _resueltos_por_estudiante_ep(ids, [self.estudiante.id])

        comision_b = Comision.objects.get(pk=self.comision_b.pk)
        avance_b = _calcular_avance_promedio(comision_b, [self.estudiante.id], resueltos)
        self.assertEqual(avance_b, 0.0, 'B no debería contar lo resuelto en A')

    def test_aprobado_en_A_no_debe_contar_en_B(self):
        """Mismo escenario sobre la función preexistente, para acotar el alcance."""
        from docentes.views import _aprobados_por_estudiante_ep, _calcular_avance_promedio
        Intento.objects.create(
            estudiante=self.estudiante, ejercicio_practica=self.ep,
            practica_comision=self.pc_a, cohorte=self.cohorte,
            respuesta_raw='p', es_correcto=True, aprobado_docente=True,
        )
        ids = [self.comision_a.id, self.comision_b.id]
        aprobados = _aprobados_por_estudiante_ep(ids, [self.estudiante.id])

        comision_b = Comision.objects.get(pk=self.comision_b.pk)
        avance_b = _calcular_avance_promedio(comision_b, [self.estudiante.id], aprobados)
        self.assertEqual(avance_b, 0.0, 'B no debería contar lo aprobado en A')


class PracticaComisionFormModoTests(TestCase):
    def setUp(self):
        self.docente = _u('doc-form-modo', es_docente=True)
        self.comision = Comision.objects.create(nombre='IPC Form Modo')
        self.comision.docentes.add(self.docente)

    def _data(self, **extra):
        data = {'comision': self.comision.id, 'orden': 1}
        data.update(extra)
        return data

    def test_sin_marcar_queda_en_libre(self):
        from docentes.forms import PracticaComisionForm
        form = PracticaComisionForm(data=self._data(), usuario=self.docente)

        self.assertTrue(form.is_valid(), form.errors)
        self.assertFalse(form.cleaned_data['desbloqueo_secuencial'])

    def test_marcado_queda_en_secuencial(self):
        from docentes.forms import PracticaComisionForm
        form = PracticaComisionForm(
            data=self._data(desbloqueo_secuencial='on'), usuario=self.docente,
        )

        self.assertTrue(form.is_valid(), form.errors)
        self.assertTrue(form.cleaned_data['desbloqueo_secuencial'])


class PanelModoLibreTests(TestCase):
    def setUp(self):
        self.docente = _u('doc-panel-libre', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Panel Libre')
        self.comision.docentes.add(self.docente)
        self.estudiante = _u('est-panel-libre')
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte)
        self.practica = Practica.objects.create(titulo='P-panel', creada_por=self.docente)
        self.pc = PracticaComision.objects.create(
            practica=self.practica, comision=self.comision, orden=1,
            desbloqueo_secuencial=False,
        )
        ej = Ejercicio.objects.create(
            enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        EjercicioPractica.objects.create(practica=self.practica, ejercicio=ej, orden=1)

    def test_no_muestra_ejercicio_actual_en_practicas_libres(self):
        self.client.login(username='doc-panel-libre', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Ej. 1')
        # La práctica tiene un ejercicio: rotularla "sin ejercicios" sería falso.
        self.assertNotContains(response, 'sin ejercicios')
        detalle = response.context['progreso_estudiantes'][0]['practicas_detalle'][0]
        self.assertFalse(detalle['desbloqueo_secuencial'])

    def test_sin_ejercicios_se_reserva_para_practicas_vacias(self):
        vacia = Practica.objects.create(titulo='P-vacia', creada_por=self.docente)
        PracticaComision.objects.create(
            practica=vacia, comision=self.comision, orden=2,
            desbloqueo_secuencial=False,
        )

        self.client.login(username='doc-panel-libre', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertContains(response, 'sin ejercicios')

    def test_secuencial_terminada_sin_revisar_no_dice_sin_ejercicios(self):
        """Progreso completo pero sin aprobación docente: ejercicio_actual es None."""
        self.pc.desbloqueo_secuencial = True
        self.pc.save(update_fields=['desbloqueo_secuencial'])
        Progreso.objects.create(
            estudiante=self.estudiante, practica_comision=self.pc, cohorte=self.cohorte,
            ejercicio_practica_actual=None,
        )

        self.client.login(username='doc-panel-libre', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertNotContains(response, 'sin ejercicios')

    def test_si_lo_muestra_en_practicas_secuenciales(self):
        self.pc.desbloqueo_secuencial = True
        self.pc.save(update_fields=['desbloqueo_secuencial'])

        self.client.login(username='doc-panel-libre', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertContains(response, 'Ej. 1')


class ProgresoDocenteAisladoPorComisionTests(TestCase):
    """El panel docente lee el Progreso de su comisión, no el de la otra."""

    def setUp(self):
        self.docente = _u('doc-panel-aisl', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.estudiante = _u('est-panel-aisl')
        self.comision_a = Comision.objects.create(nombre='C-panel-A')
        self.comision_b = Comision.objects.create(nombre='C-panel-B')
        self.comision_a.docentes.add(self.docente)
        self.comision_b.docentes.add(self.docente)
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision_a, cohorte=self.cohorte)
        Inscripcion.objects.create(estudiante=self.estudiante, comision=self.comision_b, cohorte=self.cohorte)

        self.practica = Practica.objects.create(titulo='P-panel', creada_por=self.docente)
        self.pc_a = PracticaComision.objects.create(
            practica=self.practica, comision=self.comision_a, orden=1,
        )
        self.pc_b = PracticaComision.objects.create(
            practica=self.practica, comision=self.comision_b, orden=1,
        )
        self.eps = []
        for i in (1, 2):
            ej = Ejercicio.objects.create(
                enunciado=f'E{i}', formula_solucion='p',
                tipo='formalizacion', creado_por=self.docente,
            )
            self.eps.append(EjercicioPractica.objects.create(
                practica=self.practica, ejercicio=ej, orden=i,
            ))
        # En A el estudiante ya va por el ejercicio 2; en B no arrancó.
        Progreso.objects.create(
            estudiante=self.estudiante, cohorte=self.cohorte,
            practica_comision=self.pc_a, ejercicio_practica_actual=self.eps[1],
        )

    def test_comision_b_no_hereda_el_ejercicio_actual_de_a(self):
        from docentes.views import _build_progreso_estudiantes
        comision_b = Comision.objects.get(pk=self.comision_b.pk)

        filas = _build_progreso_estudiantes(comision_b, set(), set())

        detalle = filas[0]['practicas_detalle'][0]
        self.assertEqual(
            detalle['ejercicio_actual'], self.eps[0],
            'B debería arrancar en el ejercicio 1, no heredar el puntero de A',
        )


class ProgresoDocenteRecursanteUsaCohorteActualTests(TestCase):
    """Code review: _build_progreso_estudiantes indexaba pc.progresos.all()
    solo por estudiante, sin filtrar por cohorte. PracticaComision no tiene
    cohorte propia (pertenece a la Comision, que persiste entre camadas), así
    que un recursante de la misma comisión tiene dos filas de Progreso para
    el mismo practica_comision, una por cohorte. El dict comprehension se
    quedaba con la última que iteraba, arbitraria, y el panel docente podía
    mostrar el ejercicio_actual de la camada equivocada."""

    def setUp(self):
        self.docente = _u('doc-recursa-panel', es_docente=True)
        self.estudiante = _u('est-recursa-panel')

        # Ids de Cohorte deliberadamente al revés del orden cronológico real
        # (igual que en ejercicios.tests.RecursanteReadPathTests): cohorte_de()
        # se basa en la fecha/orden de Inscripcion, no en el id de Cohorte.
        self.cohorte_actual = Cohorte.objects.create(anio=2030, cuatrimestre=1)
        self.cohorte_vieja = Cohorte.objects.create(anio=2020, cuatrimestre=1)

        # Quien recursa está inscripto en la camada EN CURSO: el panel filtra
        # por la cohorte activa, así que sin esto la vista mostraría —con
        # razón— una comisión vacía y el test no probaría lo que dice probar.
        # Se desactiva primero la 2026-C1 del backfill: el constraint
        # unica_cohorte_activa no admite dos activas.
        Cohorte.objects.filter(activa=True).update(activa=False)
        Cohorte.objects.filter(pk=self.cohorte_actual.pk).update(activa=True)
        self.cohorte_actual.refresh_from_db()

        self.comision = Comision.objects.create(nombre='IPC Recursa Panel')
        self.comision.docentes.add(self.docente)

        # Inscripción vieja primero, luego la actual: cohorte_de() debe
        # resolver a cohorte_actual (la inscripción más reciente).
        Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte_vieja,
        )
        Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte_actual,
        )

        self.practica = Practica.objects.create(titulo='P-recursa-panel', creada_por=self.docente)
        self.pc = PracticaComision.objects.create(practica=self.practica, comision=self.comision, orden=1)
        ej1 = Ejercicio.objects.create(
            enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        ej2 = Ejercicio.objects.create(
            enunciado='E2', formula_solucion='q', tipo='formalizacion', creado_por=self.docente,
        )
        self.ep1 = EjercicioPractica.objects.create(practica=self.practica, ejercicio=ej1, orden=1)
        self.ep2 = EjercicioPractica.objects.create(practica=self.practica, ejercicio=ej2, orden=2)

        # Progreso de la cohorte ACTUAL: en curso, apuntando a ep2. Se crea
        # primero para que, si el bug estuviera arreglado por casualidad de
        # orden, el test no lo tapara.
        self.progreso_actual = Progreso.objects.create(
            estudiante=self.estudiante, practica_comision=self.pc, cohorte=self.cohorte_actual,
            ejercicio_practica_actual=self.ep2,
        )
        # Progreso de la cohorte VIEJA: completa (ejercicio_practica_actual
        # None). Se crea DESPUÉS: sin filtro de cohorte, el dict comprehension
        # de _build_progreso_estudiantes pisa la fila anterior con esta.
        Progreso.objects.create(
            estudiante=self.estudiante, practica_comision=self.pc, cohorte=self.cohorte_vieja,
            ejercicio_practica_actual=None,
        )

    def test_panel_muestra_ejercicio_actual_de_la_cohorte_del_estudiante(self):
        from docentes.views import _build_progreso_estudiantes
        comision = Comision.objects.get(pk=self.comision.pk)

        filas = _build_progreso_estudiantes(comision, set(), set())

        filas_estudiante = [f for f in filas if f['estudiante'].id == self.estudiante.id]
        self.assertTrue(filas_estudiante, 'Debería haber al menos una fila para el estudiante')
        for fila in filas_estudiante:
            detalle = fila['practicas_detalle'][0]
            self.assertEqual(
                detalle['ejercicio_actual'], self.ep2,
                'Debería mostrar el puntero de la cohorte actual (ep2), no el '
                'de la cohorte vieja (completa, None)',
            )

    def test_panel_docente_completo_muestra_ejercicio_actual_correcto(self):
        """Mismo caso a través de la vista completa (comision_detail), para
        cubrir el prefetch real usado en producción."""
        self.client.login(username='doc-recursa-panel', password='clave123')
        response = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))

        self.assertEqual(response.status_code, 200)
        filas_estudiante = [
            f for f in response.context['progreso_estudiantes']
            if f['estudiante'].id == self.estudiante.id
        ]
        self.assertTrue(filas_estudiante)
        for fila in filas_estudiante:
            detalle = fila['practicas_detalle'][0]
            self.assertEqual(detalle['ejercicio_actual'], self.ep2)


class CohorteCreateTests(TestCase):
    def setUp(self):
        from cursos.models import Cohorte, Comision
        self.c1 = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.docente = _u('doc-cohorte', es_docente=True)
        self.staff_no_super = _u('staff-cohorte', is_staff=True, is_superuser=False)
        self.superuser = _u('super-cohorte', is_staff=True, is_superuser=True)
        self.comision = Comision.objects.create(nombre='IPC Noche')
        self.comision.docentes.add(self.docente)
        self.url = reverse('docentes:cohorte_create')

    def _post(self, anio=2026, cuatrimestre=2):
        return self.client.post(self.url, {'anio': anio, 'cuatrimestre': cuatrimestre})

    def test_docente_recibe_403(self):
        from cursos.models import Cohorte
        self.client.login(username='doc-cohorte', password='clave123')
        self.assertEqual(self._post().status_code, 403)
        self.assertEqual(Cohorte.objects.count(), 1)

    def test_staff_no_superusuario_recibe_403(self):
        """La distinción que este diseño introduce: is_staff no alcanza."""
        from cursos.models import Cohorte
        self.client.login(username='staff-cohorte', password='clave123')
        self.assertEqual(self._post().status_code, 403)
        self.assertEqual(Cohorte.objects.count(), 1)

    def test_superusuario_crea_y_activa(self):
        from cursos.models import Cohorte
        self.client.login(username='super-cohorte', password='clave123')
        self._post()
        nueva = Cohorte.objects.get(anio=2026, cuatrimestre=2)
        self.c1.refresh_from_db()
        self.assertTrue(nueva.activa)
        self.assertFalse(self.c1.activa)

    def test_duplicada_no_crea_ni_cambia_activa(self):
        from cursos.models import Cohorte
        self.client.login(username='super-cohorte', password='clave123')
        self._post(anio=2026, cuatrimestre=1)
        self.c1.refresh_from_db()
        self.assertEqual(Cohorte.objects.count(), 1)
        self.assertTrue(self.c1.activa)

    def test_rechazo_no_deja_cohorte_creada(self):
        from cursos.models import Cohorte
        self.client.login(username='doc-cohorte', password='clave123')
        self._post(anio=2027, cuatrimestre=1)
        self.assertFalse(Cohorte.objects.filter(anio=2027).exists())

    def test_next_externo_no_redirige_fuera_del_sitio(self):
        """`next` viene del cliente: un host ajeno no puede volverse destino."""
        self.client.login(username='super-cohorte', password='clave123')
        resp = self.client.post(self.url, {
            'anio': 2026, 'cuatrimestre': 2, 'next': 'https://evil.example.com/phish',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], reverse('docentes:comisiones_list'))

    def test_next_interno_se_respeta(self):
        self.client.login(username='super-cohorte', password='clave123')
        destino = reverse('docentes:comision_detail', args=[self.comision.id])
        resp = self.client.post(self.url, {
            'anio': 2026, 'cuatrimestre': 2, 'next': destino,
        })
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], destino)


class SelectorCohorteTests(TestCase):
    def setUp(self):
        from cursos.models import Cohorte, Comision, Inscripcion
        self.c1 = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.c2 = Cohorte.objects.create(anio=2026, cuatrimestre=2)
        self.docente = _u('doc-sel', es_docente=True)
        self.comision = Comision.objects.create(nombre='IPC Noche')
        self.comision.docentes.add(self.docente)

        self.est_c1 = _u('est-c1')
        self.est_c2 = _u('est-c2')
        Inscripcion.objects.create(estudiante=self.est_c1, comision=self.comision, cohorte=self.c1)
        Inscripcion.objects.create(estudiante=self.est_c2, comision=self.comision, cohorte=self.c2)

        Cohorte.objects.filter(pk=self.c1.pk).update(activa=False)
        Cohorte.objects.filter(pk=self.c2.pk).update(activa=True)

        self.client.login(username='doc-sel', password='clave123')
        self.url = reverse('docentes:comision_detail', args=[self.comision.id])

    def test_default_muestra_la_cohorte_activa(self):
        resp = self.client.get(self.url)
        usernames = [e.username for e in resp.context['estudiantes']]
        self.assertEqual(usernames, ['est-c2'])

    def test_querystring_cambia_la_cohorte(self):
        resp = self.client.get(self.url, {'cohorte': self.c1.pk})
        usernames = [e.username for e in resp.context['estudiantes']]
        self.assertEqual(usernames, ['est-c1'])

    def test_practicas_no_cambian_entre_cohortes(self):
        """El requisito central: el aula conserva sus prácticas."""
        from ejercicios.models import Practica, PracticaComision
        practica = Practica.objects.create(titulo='P1')
        PracticaComision.objects.create(practica=practica, comision=self.comision, orden=1)

        activa = self.client.get(self.url)
        vieja = self.client.get(self.url, {'cohorte': self.c1.pk})

        titulos_activa = [i['pc'].practica.titulo for i in activa.context['practicas_con_form']]
        titulos_vieja = [i['pc'].practica.titulo for i in vieja.context['practicas_con_form']]
        self.assertEqual(titulos_activa, ['P1'])
        self.assertEqual(titulos_activa, titulos_vieja)

    def test_cohorte_inexistente_cae_a_la_activa(self):
        resp = self.client.get(self.url, {'cohorte': 99999})
        self.assertEqual(resp.context['cohorte_actual'], self.c2)

    def test_alta_fallida_no_avisa_cohorte_cerrada(self):
        """estudiante_create re-renderiza el panel sin contexto de cohorte."""
        url = reverse('docentes:estudiante_create', args=[self.comision.id])
        resp = self.client.post(url, {'username': '', 'email': 'no-es-mail'})
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'cohorte cerrada')

    def test_la_activa_se_ofrece_aunque_no_tenga_inscriptos_aca(self):
        """Sin esto el selector mostraría una cohorte y la página otra."""
        from cursos.models import Comision
        vacia = Comision.objects.create(nombre='IPC Sin Inscriptos')
        vacia.docentes.add(self.docente)

        resp = self.client.get(
            reverse('docentes:comision_detail', args=[vacia.id])
        )
        self.assertEqual(resp.context['cohorte_actual'], self.c2)
        self.assertIn(self.c2, resp.context['cohortes_disponibles'])
        self.assertEqual(list(resp.context['estudiantes']), [])

    def test_solo_lista_cohortes_con_inscripciones_en_esta_comision(self):
        from cursos.models import Cohorte
        Cohorte.objects.create(anio=2027, cuatrimestre=1)
        resp = self.client.get(self.url)
        self.assertEqual(
            [c.pk for c in resp.context['cohortes_disponibles']],
            [self.c2.pk, self.c1.pk],
        )

    def test_parciales_filtran_por_cohorte(self):
        from cursos.models import Parcial
        from datetime import date
        Parcial.objects.create(
            comision=self.comision, cohorte=self.c1, nombre='P1 vieja',
            fecha=date(2026, 5, 1), puntaje_total=10,
        )
        activa = self.client.get(self.url)
        vieja = self.client.get(self.url, {'cohorte': self.c1.pk})
        self.assertEqual(list(activa.context['parciales']), [])
        self.assertEqual([p.nombre for p in vieja.context['parciales']], ['P1 vieja'])


class SesionCohorteTests(TestCase):
    """La selección de cohorte pasó de ser un parámetro por comisión a un
    "modo de vista" global, persistido en sesión (ver
    _resolver_cohorte_actual en docentes/views.py). Estos tests cubren esa
    persistencia entre pantallas, el override por querystring y la caída a
    la activa cuando la cohorte de sesión ya no existe.
    """

    def setUp(self):
        from cursos.models import Cohorte, Comision, Inscripcion
        self.c1 = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.c2 = Cohorte.objects.create(anio=2026, cuatrimestre=2)
        self.docente = _u('doc-sesion', es_docente=True)
        self.comision = Comision.objects.create(nombre='IPC Sesion')
        self.comision.docentes.add(self.docente)

        self.est_c1 = _u('est-sesion-c1')
        self.est_c2 = _u('est-sesion-c2')
        Inscripcion.objects.create(estudiante=self.est_c1, comision=self.comision, cohorte=self.c1)
        Inscripcion.objects.create(estudiante=self.est_c2, comision=self.comision, cohorte=self.c2)

        Cohorte.objects.filter(pk=self.c1.pk).update(activa=False)
        Cohorte.objects.filter(pk=self.c2.pk).update(activa=True)

        self.client.login(username='doc-sesion', password='clave123')
        self.url_home = reverse('docentes:comisiones_list')
        self.url_detail = reverse('docentes:comision_detail', args=[self.comision.id])

    def test_seleccion_en_el_home_persiste_al_entrar_a_una_comision(self):
        """El selector vive en el home; comision_detail la respeta sin
        necesidad de repetir el parámetro."""
        self.client.get(self.url_home, {'cohorte': self.c1.pk})

        resp = self.client.get(self.url_detail)  # sin ?cohorte=

        self.assertEqual(resp.context['cohorte_actual'], self.c1)

    def test_seleccion_persiste_de_vuelta_en_el_home(self):
        """También en la otra dirección: elegida desde una comisión, el home
        la sigue mostrando (mismo mecanismo, misma clave de sesión)."""
        self.client.get(self.url_detail, {'cohorte': self.c1.pk})

        resp = self.client.get(self.url_home)  # sin ?cohorte=

        self.assertEqual(resp.context['cohorte_actual'], self.c1)

    def test_querystring_actualiza_la_sesion_para_el_proximo_request(self):
        self.client.get(self.url_detail, {'cohorte': self.c1.pk})

        resp = self.client.get(self.url_detail)  # ya sin el parámetro

        self.assertEqual(resp.context['cohorte_actual'], self.c1)

    def test_querystring_explicito_sigue_pudiendo_volver_a_la_otra(self):
        """El override sigue funcionando en cualquier sentido, no solo una vez."""
        self.client.get(self.url_detail, {'cohorte': self.c1.pk})
        self.client.get(self.url_detail, {'cohorte': self.c2.pk})

        resp = self.client.get(self.url_detail)

        self.assertEqual(resp.context['cohorte_actual'], self.c2)

    def test_cohorte_de_sesion_borrada_cae_a_la_activa_sin_romper(self):
        """Si la cohorte fijada en sesión se borró, no debe romper: cae a la
        activa como si no hubiera sesión."""
        from cursos.models import Cohorte
        temporal = Cohorte.objects.create(anio=2099, cuatrimestre=1)
        session = self.client.session
        session['cohorte_id'] = temporal.pk
        session.save()
        temporal.delete()  # sin inscripciones: PROTECT no la bloquea

        resp = self.client.get(self.url_detail)

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['cohorte_actual'], self.c2)

    def test_cohorte_de_sesion_borrada_no_deja_referencia_colgada(self):
        """Contraparte: tras la caída a la activa, la sesión ya no apunta a
        la cohorte borrada (si no, cada request repetiría el trabajo de
        detectar que no existe)."""
        from cursos.models import Cohorte
        temporal = Cohorte.objects.create(anio=2098, cuatrimestre=1)
        session = self.client.session
        session['cohorte_id'] = temporal.pk
        session.save()
        temporal.delete()

        self.client.get(self.url_detail)

        self.assertNotIn('cohorte_id', self.client.session)

    def test_crear_cohorte_limpia_la_seleccion_pinneada(self):
        """Crear una cohorte (y activarla) es la acción que debería "saltar"
        la vista a la nueva por default, en vez de seguir mostrando la que
        estaba pinneada en sesión (que puede ser justo la que se cerró).

        Todo en la sesión del mismo usuario superusuario: hacer login() como
        otro usuario en el medio limpiaría la sesión por su cuenta (Django
        flushea la sesión al cambiar de usuario autenticado) y el test
        pasaría sin ejercitar el pop de verdad.
        """
        from cursos.models import Cohorte
        superuser = _u('super-sesion', is_staff=True, is_superuser=True)
        self.client.login(username='super-sesion', password='clave123')

        self.client.get(self.url_detail, {'cohorte': self.c1.pk})  # pinnea c1
        self.assertEqual(self.client.session.get('cohorte_id'), self.c1.pk)

        self.client.post(reverse('docentes:cohorte_create'), {'anio': 2027, 'cuatrimestre': 1})

        nueva = Cohorte.objects.get(anio=2027, cuatrimestre=1)
        self.assertTrue(nueva.activa)
        self.assertNotIn('cohorte_id', self.client.session)


class BloqueoCohorteCerradaTests(TestCase):
    def setUp(self):
        from cursos.models import Cohorte, Comision, Inscripcion
        self.c1 = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.c2 = Cohorte.objects.create(anio=2026, cuatrimestre=2)
        Cohorte.objects.filter(pk=self.c1.pk).update(activa=False)
        Cohorte.objects.filter(pk=self.c2.pk).update(activa=True)

        self.docente = _u('doc-bloqueo', es_docente=True)
        self.comision = Comision.objects.create(nombre='IPC Noche')
        self.comision.docentes.add(self.docente)
        # Un estudiante de la camada cerrada, para que `_resolver_cohorte`
        # ofrezca c1 en esta comisión (si no, cae a la activa en silencio y
        # el ?cohorte=c1 de los tests de abajo no llegaría a probar la guarda).
        self.estudiante_viejo = _u('est-viejo-bloqueo')
        Inscripcion.objects.create(
            estudiante=self.estudiante_viejo, comision=self.comision, cohorte=self.c1,
        )
        self.client.login(username='doc-bloqueo', password='clave123')

    def test_panel_de_cohorte_cerrada_no_renderiza_los_forms(self):
        """El otro requisito: no basta con que el POST rechace."""
        url = reverse('docentes:comision_detail', args=[self.comision.id])
        resp = self.client.get(url, {'cohorte': self.c1.pk})
        self.assertNotContains(resp, 'Alta individual')
        self.assertNotContains(resp, 'Importación masiva')
        self.assertContains(resp, 'cohorte cerrada')

    def test_panel_de_cohorte_activa_si_renderiza_los_forms(self):
        url = reverse('docentes:comision_detail', args=[self.comision.id])
        resp = self.client.get(url)
        self.assertContains(resp, 'Alta individual')
        self.assertContains(resp, 'Importación masiva')

    def test_parcial_create_en_get_no_muestra_el_form_de_cohorte_cerrada(self):
        """parcial_form.html es otra plantilla y no sabe de cohortes."""
        url = reverse('docentes:parcial_create', args=[self.comision.id])
        resp = self.client.get(f'{url}?cohorte={self.c1.pk}')
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(
            resp['Location'],
            reverse('docentes:comision_detail', args=[self.comision.id]),
        )

    def test_notas_precreadas_solo_para_la_camada_del_parcial(self):
        """Un parcial nuevo no debe traer notas de camadas anteriores."""
        from cursos.models import Inscripcion, NotaParcial, Parcial

        est_nuevo = _u('est-nuevo-bloqueo')
        Inscripcion.objects.create(
            estudiante=est_nuevo, comision=self.comision, cohorte=self.c2,
        )

        url = reverse('docentes:parcial_create', args=[self.comision.id])
        self.client.post(url, {
            'nombre': 'Parcial 1', 'fecha': '2026-09-01', 'puntaje_total': '10',
            'umbral_aprobacion': '4', 'umbral_promocion': '7',
        })

        parcial = Parcial.objects.get(nombre='Parcial 1')
        self.assertEqual(parcial.cohorte, self.c2)
        estudiantes = set(
            NotaParcial.objects.filter(parcial=parcial)
            .values_list('estudiante_id', flat=True)
        )
        self.assertEqual(estudiantes, {est_nuevo.id})

    def test_alta_en_cohorte_cerrada_devuelve_403(self):
        url = reverse('docentes:estudiante_create', args=[self.comision.id])
        resp = self.client.post(f'{url}?cohorte={self.c1.pk}', {
            'username': 'nuevo', 'email': 'n@e.com', 'password': 'clave123',
            'first_name': 'N', 'last_name': 'E',
        })
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(Usuario.objects.filter(username='nuevo').exists())

    def test_alta_en_cohorte_activa_funciona_e_inscribe_en_ella(self):
        from cursos.models import Inscripcion
        url = reverse('docentes:estudiante_create', args=[self.comision.id])
        self.client.post(url, {
            'username': 'nuevo2', 'email': 'n2@e.com', 'password': 'clave123',
            'first_name': 'N', 'last_name': 'E',
        })
        inscripcion = Inscripcion.objects.get(estudiante__username='nuevo2')
        self.assertEqual(inscripcion.cohorte, self.c2)

    def test_parcial_en_cohorte_cerrada_devuelve_403(self):
        from cursos.models import Parcial
        url = reverse('docentes:parcial_create', args=[self.comision.id])
        resp = self.client.post(f'{url}?cohorte={self.c1.pk}', {
            'nombre': 'P1', 'fecha': '2026-05-01', 'puntaje_total': '10',
            'umbral_aprobacion': '4', 'umbral_promocion': '7',
        })
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(Parcial.objects.filter(nombre='P1').exists())

    def test_importacion_en_cohorte_cerrada_devuelve_403(self):
        url = reverse('docentes:estudiantes_importar', args=[self.comision.id])
        resp = self.client.post(f'{url}?cohorte={self.c1.pk}', {})
        self.assertEqual(resp.status_code, 403)


class CacheAnaliticasPorCohorteTests(TestCase):
    """El caché no debe servir los números de una camada dentro de otra."""

    def setUp(self):
        from django.core.cache import cache
        cache.clear()
        from cursos.models import Cohorte, Comision, Inscripcion
        self.c1 = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.c2 = Cohorte.objects.create(anio=2026, cuatrimestre=2)
        Cohorte.objects.filter(pk=self.c1.pk).update(activa=False)
        Cohorte.objects.filter(pk=self.c2.pk).update(activa=True)

        self.docente = _u('doc-cache', es_docente=True)
        self.comision = Comision.objects.create(nombre='IPC Noche')
        self.comision.docentes.add(self.docente)
        self.est_c1 = _u('cache-c1')
        Inscripcion.objects.create(estudiante=self.est_c1, comision=self.comision, cohorte=self.c1)
        self.client.login(username='doc-cache', password='clave123')

    def test_cohortes_distintas_no_comparten_entrada(self):
        url = reverse('docentes:comision_detail', args=[self.comision.id])
        activa = self.client.get(url)
        vieja = self.client.get(url, {'cohorte': self.c1.pk})
        self.assertEqual(list(activa.context['silencio']), [])
        self.assertEqual([f['username'] for f in vieja.context['silencio']], ['cache-c1'])

    def test_bulk_aprobar_invalida_las_dos_camadas(self):
        """El bulk puede tocar intentos de varias cohortes de una vez.

        Los pares (comisión, cohorte) se calculan ANTES del update: después el
        queryset ya no matchea nada y no habría nada que invalidar.
        """
        from django.core.cache import cache
        from cursos.models import Inscripcion
        from ejercicios.models import (
            Ejercicio, EjercicioPractica, Intento, Practica, PracticaComision,
        )

        est_c2 = _u('cache-c2')
        Inscripcion.objects.create(
            estudiante=est_c2, comision=self.comision, cohorte=self.c2,
        )
        practica = Practica.objects.create(titulo='P-cache', creada_por=self.docente)
        ejercicio = Ejercicio.objects.create(
            enunciado='E', formula_solucion='p', tipo='formalizacion',
            creado_por=self.docente,
        )
        ep = EjercicioPractica.objects.create(
            practica=practica, ejercicio=ejercicio, orden=1,
        )
        pc = PracticaComision.objects.create(
            practica=practica, comision=self.comision, orden=1,
        )
        for estudiante, cohorte in ((self.est_c1, self.c1), (est_c2, self.c2)):
            Intento.objects.create(
                estudiante=estudiante, ejercicio_practica=ep, practica_comision=pc,
                cohorte=cohorte, respuesta_raw='p', es_correcto=True,
            )

        # Calentar el caché de las dos camadas.
        url = reverse('docentes:comision_detail', args=[self.comision.id])
        self.client.get(url)
        self.client.get(url, {'cohorte': self.c1.pk})
        clave_c1 = f'analiticas_{self.comision.id}_{self.c1.pk}'
        clave_c2 = f'analiticas_{self.comision.id}_{self.c2.pk}'
        self.assertIsNotNone(cache.get(clave_c1))
        self.assertIsNotNone(cache.get(clave_c2))

        self.client.post(reverse('docentes:intentos_bulk_aprobar'), {
            'comision_id': str(self.comision.id),
        })

        self.assertIsNone(cache.get(clave_c1), 'quedó servida la camada vieja')
        self.assertIsNone(cache.get(clave_c2), 'quedó servida la camada actual')


class BadgeCohorteCorreccionTests(TestCase):
    """El panel de correcciones puede cruzar todas las cohortes a propósito:
    hay que poder corregirle a quien rinde un final de una camada anterior.
    El badge es lo que permite saber de qué camada es cada intento.

    Desde que el filtro de cohorte defaultea a la activa, ver esa camada
    vieja por default dejó de pasar: hace falta pedir explícitamente "todas
    las camadas" (``?cohorte_id=todas``). Ese es justamente el caso de uso de
    los finales que motiva que la opción exista y sea visible.
    """

    def setUp(self):
        from cursos.models import Cohorte, Comision, Inscripcion
        from ejercicios.models import (
            Ejercicio, EjercicioPractica, Intento, Practica, PracticaComision,
        )

        self.c1 = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.c2 = Cohorte.objects.create(anio=2026, cuatrimestre=2)
        Cohorte.objects.filter(pk=self.c1.pk).update(activa=False)
        Cohorte.objects.filter(pk=self.c2.pk).update(activa=True)

        self.docente = _u('doc-badge', es_docente=True)
        self.comision = Comision.objects.create(nombre='IPC Noche')
        self.comision.docentes.add(self.docente)

        self.est_viejo = _u('badge-viejo')
        Inscripcion.objects.create(estudiante=self.est_viejo, comision=self.comision, cohorte=self.c1)

        practica = Practica.objects.create(titulo='P-badge', creada_por=self.docente)
        ejercicio = Ejercicio.objects.create(
            enunciado='E', formula_solucion='p', tipo='formalizacion',
            creado_por=self.docente,
        )
        ep = EjercicioPractica.objects.create(
            practica=practica, ejercicio=ejercicio, orden=1,
        )
        pc = PracticaComision.objects.create(
            practica=practica, comision=self.comision, orden=1,
        )
        Intento.objects.create(
            estudiante=self.est_viejo, ejercicio_practica=ep, practica_comision=pc,
            cohorte=self.c1, respuesta_raw='p', es_correcto=True,
        )

        self.client.login(username='doc-badge', password='clave123')
        self.url = reverse('docentes:correccion_pendiente')

    def test_default_a_la_activa_no_muestra_la_camada_vieja(self):
        resp = self.client.get(self.url)
        self.assertNotContains(resp, 'badge-viejo')

    def test_todas_las_camadas_muestra_la_cohorte_de_cada_intento(self):
        resp = self.client.get(self.url, {'cohorte_id': 'todas'})

        # El intento es de una camada cerrada y con "todas" aparece: el
        # docente tiene que poder corregirlo (caso de uso: finales).
        self.assertContains(resp, 'badge-viejo')
        self.assertContains(resp, str(self.c1))

    def test_filtro_explicito_por_la_camada_vieja_tambien_la_muestra(self):
        resp = self.client.get(self.url, {'cohorte_id': self.c1.pk})
        self.assertContains(resp, 'badge-viejo')

    def test_opcion_todas_las_camadas_esta_en_el_desplegable(self):
        """La salida tiene que estar a mano: visible incluso viendo el default."""
        resp = self.client.get(self.url)
        self.assertContains(resp, 'Todas las camadas')


class EstudianteDetailPorCohorteTests(TestCase):
    """La vista docente comparte el helper, asi que tambien separa camadas."""

    def setUp(self):
        self.docente = _u('doc-detail-coh', es_docente=True)
        self.estudiante = _u('est-detail-coh')

        self.cohorte_actual = Cohorte.objects.create(anio=2030, cuatrimestre=1)
        self.cohorte_vieja = Cohorte.objects.create(anio=2020, cuatrimestre=1)

        self.comision = Comision.objects.create(nombre='IPC Detail Coh')
        self.comision.docentes.add(self.docente)
        Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte_vieja,
        )
        Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte_actual,
        )

        practica = Practica.objects.create(titulo='P-detail-coh', creada_por=self.docente)
        self.pc = PracticaComision.objects.create(
            practica=practica, comision=self.comision, orden=1,
        )
        ejercicio = Ejercicio.objects.create(
            enunciado='E1', formula_solucion='p', tipo='formalizacion',
            creado_por=self.docente,
        )
        self.ep = EjercicioPractica.objects.create(
            practica=practica, ejercicio=ejercicio, orden=1,
        )
        Intento.objects.create(
            estudiante=self.estudiante, ejercicio_practica=self.ep,
            practica_comision=self.pc, cohorte=self.cohorte_vieja,
            respuesta_raw='p', es_correcto=True,
        )
        Intento.objects.create(
            estudiante=self.estudiante, ejercicio_practica=self.ep,
            practica_comision=self.pc, cohorte=self.cohorte_actual,
            respuesta_raw='q', es_correcto=False,
        )
        self.client.login(username='doc-detail-coh', password='clave123')

    def test_el_panel_docente_separa_las_camadas(self):
        resp = self.client.get(
            reverse('docentes:estudiante_detail', args=[self.estudiante.id]),
        )

        self.assertEqual(resp.status_code, 200)
        secciones = resp.context['historial_por_cohorte']
        self.assertEqual(
            [s['cohorte'].pk for s in secciones],
            [self.cohorte_actual.pk, self.cohorte_vieja.pk],
        )
        self.assertEqual([len(s['grupos']) for s in secciones], [1, 1])


class NextSeguroEnPanelDocenteTests(TestCase):
    """`next` llega del cliente: ninguna vista puede redirigir a otro host.

    Cada vista se prueba dos veces: con un `next` de host ajeno (debe caer al
    destino por defecto, ignorando el parámetro) y con uno interno (debe
    respetarse, sin cambiar el comportamiento legítimo). Mismo criterio que
    :class:`CohorteCreateTests`.
    """

    EXTERNO = 'https://evil.example.com/phish'

    def setUp(self):
        self.docente = _u('doc-next', es_docente=True)
        self.cohorte = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.comision = Comision.objects.create(nombre='IPC Next')
        self.comision.docentes.add(self.docente)

        self.estudiante = _u('alu-next')
        Inscripcion.objects.create(
            estudiante=self.estudiante, comision=self.comision, cohorte=self.cohorte
        )

        self.client.login(username='doc-next', password='clave123')

        # Destino interno de control: distinto del destino por defecto de cada
        # vista, para que "se respeta" no se confunda con "cayó al default".
        self.interno = reverse('docentes:comision_detail', args=[self.comision.id])
        self._orden = 0

    # ── Fábricas de fixtures ─────────────────────────────────────────
    def _practica_comision(self, titulo='Práctica Next'):
        self._orden += 1
        practica = Practica.objects.create(titulo=titulo, creada_por=self.docente)
        return PracticaComision.objects.create(
            practica=practica, comision=self.comision, orden=self._orden
        )

    def _ejercicio(self, enunciado='E next'):
        return Ejercicio.objects.create(
            enunciado=enunciado,
            formula_solucion='p',
            tipo='formalizacion',
            creado_por=self.docente,
        )

    def _intento(self):
        pc = self._practica_comision('Práctica Intento')
        ep = EjercicioPractica.objects.create(
            practica=pc.practica, ejercicio=self._ejercicio(), orden=1
        )
        return Intento.objects.create(
            estudiante=self.estudiante,
            ejercicio_practica=ep,
            practica_comision=pc,
            cohorte=self.cohorte,
            respuesta_raw='p',
            es_correcto=True,
            aprobado_docente=None,
            comentario_docente='',
        )

    # ── practica_create ──────────────────────────────────────────────
    def _post_practica_create(self, next_url):
        return self.client.post(reverse('docentes:practica_create'), {
            'titulo': 'Práctica Creada',
            'comision': self.comision.id,
            'orden': 1,
            'next': next_url,
        })

    def test_practica_create_next_externo_cae_al_default(self):
        resp = self._post_practica_create(self.EXTERNO)
        practica = Practica.objects.get(titulo='Práctica Creada')
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(
            resp['Location'],
            reverse('docentes:practica_detail', args=[practica.id]),
        )

    def test_practica_create_next_interno_se_respeta(self):
        resp = self._post_practica_create(self.interno)
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], self.interno)

    # ── practica_delete ──────────────────────────────────────────────
    def test_practica_delete_next_externo_cae_al_default(self):
        pc = self._practica_comision('Práctica Borrada')
        resp = self.client.post(
            reverse('docentes:practica_delete', args=[pc.id]),
            {'next': self.EXTERNO},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(
            resp['Location'],
            reverse('docentes:comision_detail', args=[self.comision.id]),
        )

    def test_practica_delete_next_interno_se_respeta(self):
        pc = self._practica_comision('Práctica Borrada')
        destino = reverse('docentes:comisiones_list')
        resp = self.client.post(
            reverse('docentes:practica_delete', args=[pc.id]),
            {'next': destino},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], destino)

    # ── ejercicio_delete ─────────────────────────────────────────────
    def test_ejercicio_delete_next_externo_cae_al_default(self):
        ejercicio = self._ejercicio('E borrado')
        resp = self.client.post(
            reverse('docentes:ejercicio_delete', args=[ejercicio.id]),
            {'next': self.EXTERNO},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], reverse('docentes:comisiones_list'))

    def test_ejercicio_delete_next_interno_se_respeta(self):
        ejercicio = self._ejercicio('E borrado')
        resp = self.client.post(
            reverse('docentes:ejercicio_delete', args=[ejercicio.id]),
            {'next': self.interno},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], self.interno)

    # ── intento_comentario_update ────────────────────────────────────
    def test_intento_comentario_next_externo_cae_al_default(self):
        intento = self._intento()
        resp = self.client.post(
            reverse('docentes:intento_comentario_update', args=[intento.id]),
            {'comentario_docente': 'Revisar la polaridad.', 'next': self.EXTERNO},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(
            resp['Location'],
            reverse('docentes:estudiante_detail', args=[self.estudiante.id]),
        )

    def test_intento_comentario_next_interno_se_respeta(self):
        intento = self._intento()
        resp = self.client.post(
            reverse('docentes:intento_comentario_update', args=[intento.id]),
            {'comentario_docente': 'Revisar la polaridad.', 'next': self.interno},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], self.interno)

    # ── intento_aprobacion_update: rechazo sin comentario ────────────
    def test_rechazo_sin_comentario_next_externo_cae_al_default(self):
        intento = self._intento()
        resp = self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[intento.id]),
            {'accion': 'rechazar', 'comentario_docente': '', 'next': self.EXTERNO},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(
            resp['Location'],
            reverse('docentes:estudiante_detail', args=[self.estudiante.id]),
        )

    def test_rechazo_sin_comentario_next_interno_se_respeta(self):
        intento = self._intento()
        resp = self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[intento.id]),
            {'accion': 'rechazar', 'comentario_docente': '', 'next': self.interno},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], self.interno)

    # ── intento_aprobacion_update: camino normal ─────────────────────
    def test_aprobacion_next_externo_cae_al_default(self):
        intento = self._intento()
        resp = self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[intento.id]),
            {'accion': 'aprobar', 'comentario_docente': '', 'next': self.EXTERNO},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(
            resp['Location'],
            reverse('docentes:estudiante_detail', args=[self.estudiante.id]),
        )

    def test_aprobacion_next_interno_se_respeta(self):
        intento = self._intento()
        resp = self.client.post(
            reverse('docentes:intento_aprobacion_update', args=[intento.id]),
            {'accion': 'aprobar', 'comentario_docente': '', 'next': self.interno},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp['Location'], self.interno)


class ReinscripcionEstudianteTests(TestCase):
    """Último hallazgo abierto del PR #189: no había ninguna forma de crear un
    recursante desde el panel. Cubre los dos caminos de reinscripción:

    1. `estudiante_reinscribir_recursante`: un `<select>` con quienes tienen
       historial en otra camada de esta misma comisión.
    2. `estudiante_reinscribir_pase`: búsqueda por usuario/email EXACTO de una
       cuenta de otra comisión, con pantalla de confirmación intermedia.

    Y la propiedad que hace que el modelo de cohortes tenga sentido: la
    reinscripción no debe tocar el progreso ni los intentos de la camada
    anterior del recursante.
    """

    def setUp(self):
        self.docente = _u('doc-reinscribe', es_docente=True)
        self.docente_otro = _u('doc-reinscribe-otro', es_docente=True)

        # 2026-C1 es la activa del backfill.
        self.cohorte_actual = Cohorte.objects.get(anio=2026, cuatrimestre=1)
        self.cohorte_vieja = Cohorte.objects.create(anio=2020, cuatrimestre=1)

        self.comision = Comision.objects.create(nombre='IPC Reinscripcion')
        self.comision.docentes.add(self.docente)

        self.comision_otra = Comision.objects.create(nombre='Otra comisión')
        self.comision_otra.docentes.add(self.docente_otro)

        # Recursante: cursó esta comisión en la camada vieja, no en la actual.
        self.recursante = _u('recursa-reinscribe', first_name='Ada', last_name='Lovelace')
        Inscripcion.objects.create(
            estudiante=self.recursante, comision=self.comision, cohorte=self.cohorte_vieja,
        )

        # Ya reinscripto en ambas camadas: no debe volver a ofrecerse.
        self.ya_reinscripto = _u('ya-reinscripto', first_name='Grace', last_name='Hopper')
        Inscripcion.objects.create(
            estudiante=self.ya_reinscripto, comision=self.comision, cohorte=self.cohorte_vieja,
        )
        Inscripcion.objects.create(
            estudiante=self.ya_reinscripto, comision=self.comision, cohorte=self.cohorte_actual,
        )

        # Pase: cuenta de otra comisión, sin ninguna Inscripcion en self.comision.
        self.pase = _u(
            'pase-otra-comision', email='pase@example.com',
            first_name='Katherine', last_name='Johnson',
        )
        Inscripcion.objects.create(
            estudiante=self.pase, comision=self.comision_otra, cohorte=self.cohorte_actual,
        )

        self.docente_target = _u('doc-target-pase', es_docente=True)
        self.staff_target = _u('staff-target-pase', is_staff=True)

        self.client.login(username='doc-reinscribe', password='clave123')

    # ── Camino 1: recursante (select) ────────────────────────────────

    def test_recursantes_disponibles_incluye_al_recursante_y_excluye_al_ya_inscripto(self):
        resp = self.client.get(reverse('docentes:comision_detail', args=[self.comision.id]))
        ids = set(resp.context['recursantes_disponibles'].values_list('id', flat=True))
        self.assertIn(self.recursante.id, ids)
        self.assertNotIn(self.ya_reinscripto.id, ids)
        # El pase (sin historial en esta comisión) tampoco es un recursante acá.
        self.assertNotIn(self.pase.id, ids)

    def test_reinscribir_recursante_crea_inscripcion_en_cohorte_actual(self):
        url = reverse('docentes:estudiante_reinscribir_recursante', args=[self.comision.id])
        resp = self.client.post(url, {'estudiante_id': self.recursante.id})

        self.assertEqual(resp.status_code, 302)
        self.assertTrue(
            Inscripcion.objects.filter(
                estudiante=self.recursante, comision=self.comision, cohorte=self.cohorte_actual,
            ).exists()
        )
        # La inscripción de la camada vieja no se toca.
        self.assertTrue(
            Inscripcion.objects.filter(
                estudiante=self.recursante, comision=self.comision, cohorte=self.cohorte_vieja,
            ).exists()
        )

    def test_reinscribir_recursante_ya_inscripto_en_actual_no_se_puede_forzar(self):
        """No aparece en el select; forzar el id igual no crea una tercera fila."""
        url = reverse('docentes:estudiante_reinscribir_recursante', args=[self.comision.id])
        resp = self.client.post(url, {'estudiante_id': self.ya_reinscripto.id}, follow=True)

        self.assertContains(resp, 'Seleccioná un recursante válido de la lista.')
        self.assertEqual(
            Inscripcion.objects.filter(
                estudiante=self.ya_reinscripto, comision=self.comision,
            ).count(),
            2,
        )

    def test_reinscribir_recursante_sin_historial_en_la_comision_falla(self):
        """Alguien sin ninguna Inscripcion acá no es un recursante válido de
        este camino (le corresponde el camino 2, pase)."""
        url = reverse('docentes:estudiante_reinscribir_recursante', args=[self.comision.id])
        resp = self.client.post(url, {'estudiante_id': self.pase.id})

        self.assertEqual(resp.status_code, 302)
        self.assertFalse(
            Inscripcion.objects.filter(estudiante=self.pase, comision=self.comision).exists()
        )

    def test_reinscribir_recursante_docente_ajeno_no_puede(self):
        self.client.login(username='doc-reinscribe-otro', password='clave123')
        url = reverse('docentes:estudiante_reinscribir_recursante', args=[self.comision.id])
        resp = self.client.post(url, {'estudiante_id': self.recursante.id})

        self.assertEqual(resp.status_code, 403)
        self.assertFalse(
            Inscripcion.objects.filter(
                estudiante=self.recursante, comision=self.comision, cohorte=self.cohorte_actual,
            ).exists()
        )

    def test_reinscribir_recursante_en_cohorte_cerrada_403(self):
        url = reverse('docentes:estudiante_reinscribir_recursante', args=[self.comision.id])
        resp = self.client.post(
            f'{url}?cohorte={self.cohorte_vieja.pk}', {'estudiante_id': self.recursante.id},
        )
        self.assertEqual(resp.status_code, 403)

    def test_reinscribir_recursante_sin_cohorte_activa_403(self):
        Cohorte.objects.filter(activa=True).update(activa=False)
        url = reverse('docentes:estudiante_reinscribir_recursante', args=[self.comision.id])
        resp = self.client.post(url, {'estudiante_id': self.recursante.id})
        self.assertEqual(resp.status_code, 403)

    def test_reinscribir_recursante_no_toca_progreso_ni_intentos_de_la_camada_anterior(self):
        """La propiedad que hace que el modelo de cohortes tenga sentido:
        reinscribir en la camada actual no debe alterar ni el Progreso ni los
        Intento de la camada anterior, y no debe crear un Progreso nuevo por
        arte de magia en la camada actual."""
        practica = Practica.objects.create(titulo='P-reinscribe', creada_por=self.docente)
        pc = PracticaComision.objects.create(practica=practica, comision=self.comision, orden=1)
        ejercicio = Ejercicio.objects.create(
            enunciado='E1', formula_solucion='p', tipo='formalizacion', creado_por=self.docente,
        )
        ep = EjercicioPractica.objects.create(practica=practica, ejercicio=ejercicio, orden=1)
        intento = Intento.objects.create(
            estudiante=self.recursante, ejercicio_practica=ep, practica_comision=pc,
            cohorte=self.cohorte_vieja, respuesta_raw='p', es_correcto=True,
        )
        progreso = Progreso.objects.create(
            estudiante=self.recursante, practica_comision=pc, cohorte=self.cohorte_vieja,
            ejercicio_practica_actual=None,
        )

        url = reverse('docentes:estudiante_reinscribir_recursante', args=[self.comision.id])
        self.client.post(url, {'estudiante_id': self.recursante.id})

        intento.refresh_from_db()
        progreso.refresh_from_db()
        self.assertEqual(intento.cohorte, self.cohorte_vieja)
        self.assertEqual(progreso.cohorte, self.cohorte_vieja)
        self.assertFalse(
            Progreso.objects.filter(
                estudiante=self.recursante, practica_comision=pc, cohorte=self.cohorte_actual,
            ).exists()
        )

    # ── Camino 2: pase (búsqueda por usuario/email exacto) ───────────

    def test_pase_busca_por_username_exacto_muestra_confirmacion_sin_inscribir(self):
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'pase-otra-comision'})

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['pase_pendiente'], self.pase)
        self.assertFalse(
            Inscripcion.objects.filter(
                estudiante=self.pase, comision=self.comision, cohorte=self.cohorte_actual,
            ).exists()
        )

    def test_pase_busca_por_email_exacto(self):
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'pase@example.com'})

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['pase_pendiente'], self.pase)

    def test_pase_confirmar_crea_la_inscripcion(self):
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'confirmar': '1', 'estudiante_id': self.pase.id})

        self.assertEqual(resp.status_code, 302)
        self.assertTrue(
            Inscripcion.objects.filter(
                estudiante=self.pase, comision=self.comision, cohorte=self.cohorte_actual,
            ).exists()
        )
        # La inscripción original en la otra comisión no se toca.
        self.assertTrue(
            Inscripcion.objects.filter(estudiante=self.pase, comision=self.comision_otra).exists()
        )

    def test_pase_no_existe_ninguna_cuenta(self):
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'no-existe-nadie'}, follow=True)

        self.assertContains(resp, 'No existe ninguna cuenta')
        self.assertFalse(Usuario.objects.filter(username='no-existe-nadie').exists())

    def test_pase_cuenta_docente_rechazada(self):
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'doc-target-pase'}, follow=True)

        self.assertContains(resp, 'pertenece a un docente o a personal del staff')
        self.assertFalse(
            Inscripcion.objects.filter(estudiante=self.docente_target, comision=self.comision).exists()
        )

    def test_pase_cuenta_staff_rechazada(self):
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'staff-target-pase'}, follow=True)

        self.assertContains(resp, 'pertenece a un docente o a personal del staff')
        self.assertFalse(
            Inscripcion.objects.filter(estudiante=self.staff_target, comision=self.comision).exists()
        )

    def test_pase_ya_inscripto_en_cohorte_actual_no_hace_nada(self):
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'ya-reinscripto'}, follow=True)

        self.assertContains(resp, 'ya está inscripto')
        self.assertEqual(
            Inscripcion.objects.filter(estudiante=self.ya_reinscripto, comision=self.comision).count(),
            2,
        )

    def test_pase_no_hace_busqueda_parcial(self):
        """Con icontains, 'pase' matchearía 'pase-otra-comision'; con
        coincidencia exacta no matchea nada."""
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'pase'}, follow=True)

        self.assertContains(resp, 'No existe ninguna cuenta')

    def test_pase_docente_ajeno_no_puede(self):
        self.client.login(username='doc-reinscribe-otro', password='clave123')
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'pase-otra-comision'})

        self.assertEqual(resp.status_code, 403)

    def test_pase_en_cohorte_cerrada_403(self):
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(
            f'{url}?cohorte={self.cohorte_vieja.pk}', {'identificador': 'pase-otra-comision'},
        )
        self.assertEqual(resp.status_code, 403)

    def test_pase_sin_cohorte_activa_403(self):
        Cohorte.objects.filter(activa=True).update(activa=False)
        url = reverse('docentes:estudiante_reinscribir_pase', args=[self.comision.id])
        resp = self.client.post(url, {'identificador': 'pase-otra-comision'})
        self.assertEqual(resp.status_code, 403)
